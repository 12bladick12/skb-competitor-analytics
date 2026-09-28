from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from bs4 import BeautifulSoup
from openpyxl import Workbook,load_workbook

from price_monitor.adapters import ADAPTERS,parse_money
from price_monitor.exchange import COLUMNS,read_table,validate_rows,xlsx_bytes,csv_bytes
from price_monitor.models import Rule,Observation
from price_monitor.robots import Robots
from price_monitor.sources import SOURCES,validate_url
from price_monitor.storage import Store
from price_monitor.transport import SourceClient,FetchError,challenge
from price_monitor.worker import Worker
from support import TestDirectory

FIXTURES = Path(__file__).parent/"fixtures"
MANIFEST = json.loads((FIXTURES/"manifest.json").read_text(encoding="utf-8"))


def case_rule(case):
    return Rule(**{k:case[k] for k in ("source","manufacturer","article","product_url")})


class ParsingTests(unittest.TestCase):
    def test_actual_cards_all_brands_and_discontinued_model(self):
        for case in MANIFEST:
            with self.subTest(case=case["article"]):
                html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
                result=ADAPTERS[case["source"]].parse(case_rule(case),html)
                self.assertEqual(result.status,case["status"],result.detail)
                self.assertEqual(result.price,case["price"])
                self.assertEqual(result.currency,"RUB" if case["price"] else None)

    def test_no_public_price_and_request_price_for_every_source(self):
        for source in SOURCES:
            case=next(c for c in MANIFEST if c["source"]==source and c["price"])
            rule=case_rule(case)
            html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
            soup=BeautifulSoup(html,"html.parser")
            root=soup.select_one(SOURCES[source].scope)
            for node in list(root.select(SOURCES[source].price_selector)):
                node.decompose()
            with self.subTest(source=source,mode="no price"):
                result=ADAPTERS[source].parse(rule,str(soup))
                self.assertEqual(result.status,"no_price",result.detail)
                self.assertIsNone(result.price)
            area={"sensoren":".product-info__all-order","beskonta":".p-p-block-price","megak":".details-payment","teko":".catalog-detail-right","sensor":".navigation-product__priceblock"}[source]
            root.select_one(area).append("Цена по запросу")
            with self.subTest(source=source,mode="on request"):
                result=ADAPTERS[source].parse(rule,str(soup))
                self.assertEqual(result.status,"on_request",result.detail)
                self.assertIsNone(result.price)

    def test_not_found_and_wrong_model_every_source(self):
        for source in SOURCES:
            case=next(c for c in MANIFEST if c["source"]==source)
            rule=case_rule(case)
            self.assertEqual(ADAPTERS[source].parse(rule,"<h1>404</h1>",http_status=404).status,"not_found")
            html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
            result=ADAPTERS[source].parse(replace(rule,article="NOT-EXIST-987"),html)
            self.assertIn(result.status,{"identity_mismatch","needs_variant"})
            self.assertIsNone(result.price)

    def test_no_price_from_related_products(self):
        case=MANIFEST[-1]
        html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
        html += '<div class="product-box__price">5 917.00 ₽</div>'
        result=ADAPTERS["sensor"].parse(case_rule(case),html)
        self.assertEqual(result.availability,"discontinued")
        self.assertIsNone(result.price)
        self.assertIn("5 917",html)  # recommended replacement really has a price

    def test_variant_and_manufacturer_mismatch(self):
        case=next(c for c in MANIFEST if c["source"]=="beskonta")
        html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
        self.assertEqual(ADAPTERS["beskonta"].parse(replace(case_rule(case),article="BORE01"),html).status,"needs_variant")
        case=MANIFEST[0]
        html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
        self.assertEqual(ADAPTERS["sensoren"].parse(replace(case_rule(case),manufacturer="SICK"),html).status,"identity_mismatch")

    def test_broken_selector_and_ambiguous_price(self):
        case=MANIFEST[0]; rule=case_rule(case)
        self.assertEqual(ADAPTERS[rule.source].parse(rule,"<h1>SI5000</h1>").status,"parse_error")
        html=(FIXTURES/case["file"]).read_text(encoding="utf-8")
        html=html.replace('class="price"','class="price-old"')
        result=ADAPTERS[rule.source].parse(rule,html)
        self.assertIsNone(result.price)
        self.assertEqual(result.status,'parse_error')

    def test_money_does_not_mix_vat_units_or_old_prices(self):
        for text,expected in [("1\xa0708 ₽","1708.00"),("5 917.00 ₽","5917.00"),("29 530 руб. / шт","29530.00"),("28 960 ₽ НДС 22%",None),("от 500 ₽",None),("100 ₽ 90 ₽",None),("0 ₽",None)]:
            self.assertEqual(parse_money(text),expected)


class ImportExportTests(unittest.TestCase):
    def test_csv_xlsx_roundtrip_and_leading_zero(self):
        row={"source":"sensoren","manufacturer":"ifm","article":"00123","product_url":"https://sensoren.ru/product/00123/","url_template":""}
        for name,data in [("x.csv",csv_bytes([row])),("x.xlsx",xlsx_bytes([row],COLUMNS,"Задания"))]:
            rules,errors=validate_rows(read_table(data,name))
            self.assertFalse(errors)
            self.assertEqual(rules[0].article,"00123")

    def test_invalid_rows_duplicate_and_external_urls(self):
        row={"source":"sensoren","manufacturer":"ifm","article":"SI5000","product_url":"https://sensoren.ru/product/test/"}
        _,errors=validate_rows([row,row])
        self.assertEqual(len(errors),1)
        for bad in [{**row,"product_url":"http://127.0.0.1/product/a/"},{**row,"manufacturer":"BESKONTA"},{**row,"url_template":"https://sensoren.ru/product/{article}/"}]:
            self.assertTrue(validate_rows([bad])[1])

    def test_bad_headers_formulas_and_empty_file(self):
        for data in [b"",b"source;source;article\na;b;c",b"source;article\na;b"]:
            with self.assertRaises(ValueError):read_table(data,"bad.csv")
        wb=Workbook();ws=wb.active;ws.append(COLUMNS);ws.append(["sensoren","ifm","=1+1","https://sensoren.ru/product/a/",""])
        out=BytesIO();wb.save(out)
        with self.assertRaises(ValueError):read_table(out.getvalue(),"bad.xlsx")

    def test_template_and_formula_export(self):
        rows=[{"source":"megak","manufacturer":"МЕГА-К","article":"PS2 A/1","url_template":"https://mega-k.com/products/{article}"}]
        rules,errors=validate_rows(rows)
        self.assertFalse(errors)
        self.assertTrue(rules[0].url.endswith("PS2%20A%2F1"))
        payload=[{"text":"=HYPERLINK(\"https://evil\")","price":None}]
        wb=load_workbook(BytesIO(xlsx_bytes(payload)))
        self.assertEqual(wb.active['A2'].data_type,"s")
        self.assertIsNone(wb.active['B2'].value)
        self.assertIn("'=HYPERLINK",csv_bytes(payload).decode('utf-8-sig'))


class RobotsTransportTests(unittest.TestCase):
    def test_longest_match_allow_wildcard_end_and_specific_agents(self):
        policy=Robots("User-agent: *\nDisallow: /*?*\nDisallow: /private/\nAllow: /private/public$\nAllow: /*?PAGEN\n")
        for path,allowed in [("/product/a/",True),("/product/a/?q=1",False),("/product/a/?PAGEN_1=2",True),("/private/public",True),("/private/publication",False)]:
            self.assertEqual(policy.allows("https://sensoren.ru"+path),allowed)
        self.assertTrue(Robots("User-agent: *\nDisallow: /\nUser-agent: PriceMonitor\nAllow: /\n").allows("https://x/product"))

    def test_audited_robots_allow_cards_and_deny_search(self):
        for source in SOURCES:
            policy=Robots((FIXTURES/f"robots_{source}.txt").read_text(encoding="utf-8"))
            case=next(c for c in MANIFEST if c['source']==source)
            self.assertTrue(policy.allows(case['product_url']))
            denied={"sensoren":"/catalog/?q=SI5000","beskonta":"/catalog/all/?query=BORE01","megak":"/search?q=x","teko":"/search/?q=x","sensor":"/search/?query=x"}[source]
            self.assertFalse(policy.allows('https://'+SOURCES[source].host+denied))

    def client(self):
        c=SourceClient("sensoren")
        c.wait=lambda seconds:None
        return c

    def test_no_retry_for_block_or_429(self):
        for code in [401,403,429]:
            c=self.client()
            with patch.object(c,"_request",side_effect=[(200,{},"User-agent: *\nAllow: /"),(code,{"Retry-After":"120"},"blocked")]) as request:
                with self.assertRaises(FetchError) as e:c.fetch(MANIFEST[0]['product_url'])
                self.assertTrue(e.exception.stop_source)
                self.assertEqual(request.call_count,2)
            c.close()

    def test_redirect_cannot_reach_internal_address(self):
        c=self.client()
        with patch.object(c,"_request",side_effect=[(200,{},""),(302,{"Location":"http://127.0.0.1/private"},"")]) as request:
            with self.assertRaises(FetchError):c.fetch(MANIFEST[0]['product_url'])
            self.assertEqual(request.call_count,2)
        c.close()

    def test_legacy_teko_redirect_keeps_https_and_checks_robots_again(self):
        c=SourceClient('teko');c.wait=lambda _:None
        url='https://teko-com.ru/product/cp-s254r-3-pg9.html'
        with patch.object(c,'_request',side_effect=[(200,{},''),(301,{'Location':'http://teko-com.ru/catalog/product/cp-s254r-3-pg9/'},''),(200,{'Content-Type':'text/html'},'<h1>CP S254R-3</h1>')]) as req:
            final,code,_=c.fetch(url)
            self.assertEqual(final,'https://teko-com.ru/catalog/product/cp-s254r-3-pg9/')
            self.assertEqual(code,200)
            self.assertTrue(all(call.args[0].startswith('https://') for call in req.call_args_list))
        validate_url('teko',final)
        c.close()

    def test_robots_failure_and_forbidden_url(self):
        for response,expected in [((503,{},"unavailable"),"robots_unavailable"),((200,{},"User-agent: *\nDisallow: /"),"robots_denied")]:
            c=self.client()
            with patch.object(c,"_request",return_value=response) as request:
                with self.assertRaises(FetchError) as e:c.fetch(MANIFEST[0]['product_url'])
                self.assertEqual(e.exception.status,expected)
                self.assertEqual(request.call_count,1)
            c.close()

    def test_actual_form_captcha_is_not_a_challenge(self):
        self.assertFalse(challenge('<title>Датчик SI5000</title><input name="captcha_sid">'))
        self.assertTrue(challenge('<title>Just a moment...</title>'))
        for url in ['https://sensoren.ru.evil.test/product/a','https://sensoren.ru@127.0.0.1/product/a','https://sensoren.ru:8443/product/a']:
            with self.assertRaises(ValueError):validate_url('sensoren',url)


class StorageWorkerTests(unittest.TestCase):
    def setUp(self):
        self.temp=TestDirectory()
        self.db=Store(Path(self.temp.name)/'prices.db')
        self.rule=case_rule(MANIFEST[0])

    def tearDown(self):self.temp.cleanup()

    def test_concurrent_enqueue_and_durable_cancel(self):
        def enqueue():
            try:return self.db.enqueue([self.rule])
            except ValueError:return None
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:enqueue(),range(2)))
        self.assertEqual(sum(r is not None for r in results),1)
        self.db.cancel(next(r for r in results if r))
        reopened=Store(self.db.path)
        self.assertEqual(reopened.runs()[0]['state'],'cancelled')
        self.assertEqual(reopened.results()[0]['status'],'cancelled')

    def test_batched_rules_preserve_order_and_reuse_saved_rules(self):
        rules=[replace(self.rule,article=f'MODEL{i}',product_url=f'https://sensoren.ru/product/model{i}/') for i in range(205)]
        run=self.db.enqueue(rules)
        self.assertEqual([r['article'] for r in self.db.results(run)],[r.article for r in rules])
        self.db.cancel(run)
        second=self.db.enqueue(rules)
        self.assertEqual(len(self.db.saved_rules()),205)
        self.assertEqual(len(self.db.results(second)),205)

    def test_lease_recovery_and_idempotent_results(self):
        run_id=self.db.enqueue([self.rule])
        self.assertTrue(self.db.acquire('old'))
        self.assertFalse(self.db.acquire('new'))
        self.db.claim_run('old')
        job=self.db.pending(run_id)[0][0]
        self.db.processing(job,'old')
        with self.db.connect() as c:c.execute('UPDATE worker_lease SET heartbeat=0')
        self.assertTrue(self.db.acquire('new'))
        self.assertTrue(self.db.processing(job,'new'))
        result=Observation('priced',self.rule.url,price='123.45',currency='RUB')
        self.db.record(job,result,'new');self.db.record(job,result,'new')
        self.db.finish(run_id)
        self.assertEqual(len(self.db.results(run_id)),1)
        self.assertEqual(self.db.runs()[0]['state'],'completed')
        self.db.backup(Path(self.temp.name)/'backup.db')
        self.assertEqual(Store(Path(self.temp.name)/'backup.db').results()[0]['price'],'123.45')

    def test_worker_records_history_and_keeps_prior_price_on_failure(self):
        fixture=(FIXTURES/'sensoren_3.html').read_text(encoding='utf-8')
        class FakeClient:
            def __init__(self,source,**kwargs):pass
            def fetch(self,url):return url,200,fixture
            def close(self):pass
        run1=self.db.enqueue([self.rule]);Worker(self.db,FakeClient).run(once=True)
        self.assertEqual(self.db.results(run1)[0]['price'],'29530.00')
        class BlockedClient(FakeClient):
            def fetch(self,url):raise FetchError('blocked','403',403,True)
        run2=self.db.enqueue([self.rule]);Worker(self.db,BlockedClient).run(once=True)
        self.assertIsNone(self.db.results(run2)[0]['price'])
        self.assertEqual(self.db.results(run1)[0]['price'],'29530.00')
        self.assertEqual(len(self.db.results(rule_id=1)),2)

    def test_source_stop_isolated_and_user_cancel(self):
        other=case_rule(next(c for c in MANIFEST if c['source']=='teko'))
        run=self.db.enqueue([self.rule,replace(self.rule,article='SECOND',product_url='https://sensoren.ru/product/second/'),other])
        calls=[]
        class Client:
            def __init__(self,source,**kwargs):self.source=source
            def fetch(self,url):
                calls.append(self.source)
                if self.source=='sensoren':raise FetchError('rate_limited','Retry-After: 120',429,True)
                return url,404,''
            def close(self):pass
        Worker(self.db,Client).run(once=True)
        self.assertEqual(calls.count('sensoren'),1)
        self.assertEqual([r['status'] for r in self.db.results(run)],['rate_limited','source_stopped','not_found'])
        self.assertEqual(self.db.runs()[0]['state'],'completed_with_errors')

    def test_source_pause_survives_worker_restart(self):
        run=self.db.enqueue([self.rule])
        self.db.pause_source(run,'sensoren','access limited')
        class Client:
            def __init__(self,*a,**k):pass
            def fetch(self,url):raise AssertionError('must not fetch stopped source')
            def close(self):pass
        Worker(self.db,Client).run(once=True)
        self.assertEqual(self.db.results(run)[0]['status'],'source_stopped')

    def test_cancel_during_run_preserves_completed_result(self):
        second=replace(self.rule,article='SECOND',product_url='https://sensoren.ru/product/second/')
        run=self.db.enqueue([self.rule,second])
        fixture=(FIXTURES/'sensoren_3.html').read_text(encoding='utf-8')
        store=self.db
        calls=[]
        class Client:
            def __init__(self,*a,**k):pass
            def fetch(self,url):
                calls.append(url);store.cancel(run)
                return url,200,fixture
            def close(self):pass
        Worker(self.db,Client).run(once=True)
        rows=self.db.results(run)
        self.assertEqual([r['status'] for r in rows],['priced','cancelled'])
        self.assertEqual(len(calls),1)
        self.assertEqual(self.db.runs()[0]['state'],'cancelled')


if __name__ == '__main__':unittest.main()

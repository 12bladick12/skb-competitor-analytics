from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, MagicMock
from zipfile import ZipFile

from docx import Document
from cloud.coverage import monitored
from cloud.draft_rules import snapshot
from cloud.exports import render
from cloud.jobs import JobService
from test_cloud_drafts import fixture, draft_fixture


class ReleaseTests(unittest.TestCase):
    def test_large_bundle_uploads_in_bounded_parts_and_downloads_identically(self):
        import os
        from cloud.artifacts import upload_objects
        from cloud.library import read_asset
        from cloud.drive_store import StorageError
        content=os.urandom(5*1024*1024+100)
        digest=hashlib.sha256(content).hexdigest()
        store=MagicMock()
        store.query.return_value=[{'payload':{'folder_id':'test_folder_123'}}]
        store.asset_ids.return_value=set()
        store.upload_id.side_effect=lambda job,digest,allocate:allocate()
        class Drive:
            def __init__(self):self.files={};self.counter=0
            def ensure_folder(self,*args):pass
            def allocate(self):
                self.counter+=1
                return 'test_file_'+str(self.counter)
            def ensure_pack(self,ident,folder,packed,sha):
                self.assert_size=len(packed)<=5*1024*1024
                assert self.assert_size
                self.files[ident]=packed
            def download(self,ident,sha):
                value=self.files[ident]
                assert hashlib.sha256(value).hexdigest()==sha
                return value
        drive=Drive()
        assets=upload_objects(store,drive,{'import_id':'c'*64},{digest:content})
        self.assertEqual(len(assets[digest]['chunks']),2)
        repo=MagicMock()
        repo.asset.side_effect=lambda import_id,ident:assets[ident]
        self.assertEqual(read_asset(repo,drive,'c'*64,digest),content)
        assets[digest]['chunks']=[digest]
        with self.assertRaises(StorageError):read_asset(repo,drive,'c'*64,digest)

    def test_catalogue_multiple_proofs_survive_two_disposable_workers(self):
        from cloud.collection import collect
        from src.models import CompetitorConfig
        from src.verified_monitor import VerifiedMonitor
        from bs4 import BeautifulSoup
        data=fixture()
        data['event']={};data['run']={};data['period']={'2026-09':{'event_ids':[],'checks':[]}}
        company=CompetitorConfig(code='sensor',name='СЕНСОР',base_url='https://example.com',
            monitor_sources=[{'kind':'products','url':'https://example.com/products','mode':'catalogue'}])
        iteration=[0]
        def source(monitor,run_id,competitor,config,period):
            proofs=[monitor.snapshot(competitor.code,config['url']+str(i),'<html>Listing '+str(i)+'</html>') for i in range(2)]
            cards={'https://example.com/product':({'title':'Датчик','summary':'Описание '+str(iteration[0])},proofs[0])}
            added,reason=monitor.product_changes(competitor,config,cards,proofs,BeautifulSoup('<html></html>','html.parser'))
            self.assertEqual(reason,'')
            monitor.store.check(run_id,competitor.code,'products',config['url'],'success',items=added)
        def noop(*args):pass
        baselines=[];saved={}
        with patch('cloud.collection.load_competitors',return_value=[company]), \
             patch.object(VerifiedMonitor,'collect_source',source),patch.object(VerifiedMonitor,'revalidate_saved_events',noop), \
             patch.object(VerifiedMonitor,'check_event_links',noop), \
             patch.object(VerifiedMonitor,'fetch',return_value=('<html><h1>Датчик</h1><p>Подробное описание нового датчика для каталога производителя.</p></html>',200,'')):
            for number in range(2):
                iteration[0]=number
                with tempfile.TemporaryDirectory() as folder:
                    records,objects,result=collect(data,'2026-09',Path(folder),saved.__getitem__,noop,baselines)
                saved.update(objects)
                baselines=[r['payload'] for r in records if r['kind']=='collector_baseline']
                self.assertEqual(len(baselines[0]['evidence_ids']),2)
                self.assertNotIn('evidence_path',baselines[0])
                for row in records:
                    if row['kind'] in data:
                        data[row['kind']][row['key']]=row['payload']
                self.assertEqual(result['status'],'success')
        event=next(iter(data['event'].values()))
        self.assertEqual(event['date_basis'],'observed_change')
        self.assertEqual(event['status'],'confirmed')
        self.assertEqual(len(event['evidence_ids']),5)
        self.assertTrue(all(ident in saved for ident in event['evidence_ids']))

    def test_unconfigured_channels_are_outside_coverage_but_real_errors_remain(self):
        checks=[{'kind':'telegram','url':'','status':'not_configured'},
                {'kind':'telegram','url':'https://t.me/teko_pro','status':'error'},
                {'kind':'news','url':'https://example.com','status':'success','items':0}]
        self.assertEqual(monitored(checks),checks[1:])
        self.assertEqual(len(checks),3)

    def test_word_and_portable_bundle_use_saved_selection_and_inert_evidence(self):
        data=fixture()
        raw=b'<html><script>alert(1)</script><h1>evidence</h1></html>'
        digest=hashlib.sha256(raw).hexdigest()
        for event in data['event'].values():
            event['evidence_ids']=[digest]
        draft=draft_fixture(data)
        draft['items'][0].update(title='Ручной заголовок',description='Сохранённое описание')
        draft['items'][1]['included']=False
        draft['conclusions']='Вывод аналитика'
        saved=snapshot(draft,data)
        before=deepcopy(saved)
        word,bundle=render(saved,{digest:raw},data['competitor'])
        text='\n'.join(p.text for p in Document(BytesIO(word)).paragraphs)
        text+='\n'+'\n'.join(cell.text for table in Document(BytesIO(word)).tables for row in table.rows for cell in row.cells)
        self.assertIn('Ручной заголовок',text)
        self.assertIn('Сохранённое описание',text)
        self.assertIn('Вывод аналитика',text)
        self.assertNotIn(draft['items'][1]['title'],text)
        with ZipFile(BytesIO(bundle)) as archive:
            manifest=json.loads(archive.read('manifest.json'))
            self.assertEqual(len(manifest['items']),1)
            self.assertEqual(manifest['counts'],saved['counts'])
            self.assertEqual(archive.read('report.docx'),word)
            proof=manifest['items'][0]['source']['evidence_files'][0]
            self.assertEqual(archive.read(proof['original']),raw)
            self.assertNotIn(b'<script>',archive.read(proof['view']))
            self.assertIn(b'sandbox',archive.read(proof['view']))
            self.assertTrue(all(not Path(n).is_absolute() and '..' not in Path(n).parts for n in archive.namelist()))
            with tempfile.TemporaryDirectory() as folder:
                archive.extractall(folder)
                self.assertTrue((Path(folder)/proof['view']).is_file())
        self.assertEqual(saved,before)
        with self.assertRaises(ValueError):render(saved,{digest:b'changed'},data['competitor'])

    def test_job_actions_reject_viewer_before_storage(self):
        config={'access':{'viewer_emails':['reader@example.com']}}
        claims={'is_logged_in':True,'email_verified':True,'iss':'https://accounts.google.com','sub':'1','email':'reader@example.com'}
        factory=MagicMock()
        service=JobService(lambda:config,lambda:claims,store_factory=factory)
        for action in [lambda:service.collect('c'*64,'2026-09'),lambda:service.export('c'*64,'2026-09',1),lambda:service.retry('c'*64,'x')]:
            with self.assertRaises(PermissionError): action()
        factory.assert_not_called()

    def test_collection_preserves_ids_and_uses_verified_monitor_without_invented_channels(self):
        from cloud.collection import collect
        from src.models import CompetitorConfig
        from src.verified_monitor import VerifiedMonitor
        data=fixture()
        data['event']={};data['run']={};data['period']={'2026-09':{'event_ids':[],'checks':[]}}
        company=CompetitorConfig(code='sensor',name='СЕНСОР',base_url='https://example.com',
            monitor_sources=[{'kind':'news','url':'https://example.com/news','mode':'articles'}])
        page='<html><h1>Новая публикация</h1><article><p>Подробная новость производителя для проверки облачного сбора данных и сохранения доказательств.</p></article></html>'
        def source(monitor,run_id,competitor,source,period):
            path=monitor.snapshot(competitor.code,source['url'],page)
            monitor.store.record(dict(competitor_code=competitor.code,competitor_name=competitor.name,kind='news',
                url='https://example.com/article',title='Новая публикация',original_text='Подробная новость производителя для проверки и сохранения доказательств.',
                published_at='2026-09-10T00:00:00',date_basis='publication_date',
                evidence={'official_url':competitor.base_url,'source_url':source['url'],'date_evidence':'10.09.2026','snapshots':[path]}))
            monitor.store.check(run_id,competitor.code,'news',source['url'],'success',items=1)
        def noop(*args):pass
        with patch('cloud.collection.load_competitors',return_value=[company]), \
             patch.object(VerifiedMonitor,'collect_source',source),patch.object(VerifiedMonitor,'revalidate_saved_events',noop), \
             patch.object(VerifiedMonitor,'check_event_links',noop):
            saved={}
            for iteration in range(2):
                with tempfile.TemporaryDirectory() as folder:
                    records,objects,result=collect(data,'2026-09',Path(folder),saved.__getitem__,noop)
                saved.update(objects)
                for row in records:
                    if row['kind'] in data:
                        data[row['kind']][row['key']]=row['payload']
                self.assertEqual(data['period']['2026-09']['checks'][0]['items'],1)
        self.assertEqual(result['status'],'success')
        event=next(r['payload'] for r in records if r['kind']=='event')
        period=next(r['payload'] for r in records if r['kind']=='period')
        self.assertEqual(event['id'],1)
        self.assertEqual(period['event_ids'],[1])
        self.assertEqual(len(period['checks']),1)
        self.assertEqual(period['checks'][0]['kind'],'news')
        self.assertTrue(all(hashlib.sha256(v).hexdigest()==k for k,v in objects.items()))

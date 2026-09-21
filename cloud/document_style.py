"""Editorial document styling and PDF from the very same Word document content."""
from datetime import datetime
from html import escape
from pathlib import Path
import os
import shutil
from urllib.parse import urlsplit

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from src.report_builder import ReportBuilder


class EditorialStyle(ReportBuilder):
    def _setup_document(self, doc):
        super()._setup_document(doc)
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin, section.bottom_margin = Mm(22), Mm(20)
        section.left_margin = section.right_margin = Mm(21)
        section.header_distance = section.footer_distance = Mm(10)
        normal = doc.styles['Normal']
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string('263238')
        normal.paragraph_format.line_spacing = 1.2
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.widow_control = True
        for name, size in [('Title', 28), ('Subtitle', 14), ('Heading 1', 17), ('Heading 2', 12)]:
            style = doc.styles[name]
            style.font.name = 'Arial'
            style.font.size = Pt(size)
            style.font.color.rgb = self.primary
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_after = Pt(9)
        for name in ('Caption', 'Meta'):
            style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
            style.font.name = 'Arial'
            style.font.size = Pt(9)
            style.font.color.rgb = self.muted
            style.paragraph_format.space_after = Pt(5)
        header = section.header.paragraphs[0]
        header.text = 'СКБ ИНДУКЦИЯ  /  КОНКУРЕНТНАЯ АНАЛИТИКА'
        header.style = doc.styles['Caption']
        footer = section.footer.paragraphs[0]
        footer.style = doc.styles['Caption']
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Аналитическая записка   ·   ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)

    def _signature_block(self, doc):
        # Approval details belong at the end, leaving the opening for the report.
        doc.add_paragraph('Подготовил: Селезнев В.Е.  ____________________', style='Meta')
        doc.add_paragraph('Согласовал(а): Никанкина А.С.  ____________________', style='Meta')

    def _title(self, doc, competitors, period):
        doc.add_paragraph('АНАЛИТИЧЕСКАЯ ЗАПИСКА', style='Subtitle')
        doc.add_paragraph('Конкурентная среда', style='Title')
        doc.add_paragraph(period.label, style='Subtitle')
        doc.add_paragraph('Конкуренты: ' + ', '.join(c.name for c in competitors), style='Meta')
        doc.add_paragraph(f'Дата формирования: {datetime.now():%d.%m.%Y}', style='Meta')

    def _heading(self, doc, text):
        p = doc.add_paragraph(text, style='Heading 1')
        p.paragraph_format.space_before = Pt(20)
        border = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for key, value in {'val':'single', 'sz':'6', 'space':'7', 'color':'DCC3C7'}.items():
            bottom.set(qn('w:' + key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    def article(self, doc, item):
        source = item['source']
        meta = doc.add_paragraph(source['competitor_name'] + '  ·  ' + source['date_label'], style='Meta')
        meta.paragraph_format.keep_with_next = True
        meta.paragraph_format.space_before = Pt(14)
        doc.add_paragraph(item['title'], style='Heading 2')
        for line in item['description'].splitlines():
            doc.add_paragraph(line)
        target = source.get('provenance', {}).get('article_url') or source['url']
        if urlsplit(target).scheme in ('http', 'https'):
            p = doc.add_paragraph(style='Meta')
            self._add_hyperlink(p, target, 'Первоисточник ↗')

    def summary(self, doc, snapshot):
        doc.styles.add_style('Summary', 3)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Summary'
        values = [(len(snapshot['items']), 'Всего материалов'),
                  (snapshot['counts'].get('news', 0), 'Новости сайтов'),
                  (snapshot['counts'].get('telegram', 0), 'Telegram'),
                  (snapshot['counts'].get('products', 0), 'Продукция и предложения')]
        for cell, (count, label) in zip(table.rows[0].cells, values):
            cell.width = Mm(42)
            p = cell.paragraphs[0]
            p.style = doc.styles['Heading 1']
            p.add_run(str(count))
            cell.add_paragraph(label, style='Meta')
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F7F3F4')
            cell._tc.get_or_add_tcPr().append(shading)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def _table(self, doc, rows):
        super()._table(doc, rows)
        table = doc.tables[-1]
        width = Mm(168 / len(rows[0]))
        for col in table.columns:
            col.width = width
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.width = width
                if index and index % 2:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'F7F3F4')
                    cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(5)
                    paragraph.paragraph_format.space_after = Pt(5)


def document_html(doc):
    """Serialize only document text and safe hyperlinks; never interpret source HTML."""
    def inline(element):
        parts = []
        for node in element:
            if node.tag == qn('w:hyperlink'):
                relation = doc.part.rels.get(node.get(qn('r:id')))
                target = str(relation.target_ref) if relation else ''
                label = inline(node)
                if urlsplit(target).scheme in ('http', 'https'):
                    parts.append(f'<a href="{escape(target, quote=True)}">{label}</a>')
                else:
                    parts.append(label)
            elif node.tag == qn('w:t'):
                parts.append(escape(node.text or ''))
            elif node.tag == qn('w:br'):
                parts.append('<br>')
            elif node.tag == qn('w:tab'):
                parts.append(' &nbsp; ')
            elif node.tag == qn('w:r'):
                value = inline(node)
                props = node.find(qn('w:rPr'))
                if props is not None and props.find(qn('w:b')) is not None:
                    value = '<strong>' + value + '</strong>'
                parts.append(value)
        return ''.join(parts)

    def paragraph(node):
        props = node.find(qn('w:pPr'))
        style_node = props.find(qn('w:pStyle')) if props is not None else None
        style = style_node.get(qn('w:val')) if style_node is not None else ''
        tag = {'Title':'h1', 'Heading1':'h2', 'Heading2':'h3'}.get(style, 'p')
        return f'<{tag} class="{escape(style, quote=True)}">{inline(node)}</{tag}>'

    parts = []
    for node in doc.element.body:
        if node.tag == qn('w:p'):
            parts.append(paragraph(node))
        elif node.tag == qn('w:tbl'):
            properties = node.find(qn('w:tblPr'))
            table_style = properties.find(qn('w:tblStyle')) if properties is not None else None
            summary = table_style is not None and table_style.get(qn('w:val')) == 'Summary'
            rows = []
            for index, row in enumerate(node.findall(qn('w:tr'))):
                cells = [''.join(paragraph(p) for p in cell.findall(qn('w:p'))) for cell in row.findall(qn('w:tc'))]
                tag = 'th' if index == 0 and not summary else 'td'
                rows.append('<tr>' + ''.join(f'<{tag}>{cell}</{tag}>' for cell in cells) + '</tr>')
            if summary:
                parts.append('<table class="summary"><tbody>' + ''.join(rows) + '</tbody></table>')
            else:
                parts.append('<table><thead>' + rows[0] + '</thead><tbody>' + ''.join(rows[1:]) + '</tbody></table>')
    return '''<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'">
<title>Аналитическая записка по конкурентам</title><style>
* { box-sizing: border-box; } body { font: 10.5pt/1.45 Arial, sans-serif; color: #263238; margin: 0; }
h1 { font-size: 28pt; line-height: 1.12; margin: 8pt 0 14pt; color: #7a1f2b; }
h2 { font-size: 17pt; line-height: 1.2; margin: 22pt 0 12pt; padding-bottom: 7pt; border-bottom: 1px solid #dcc3c7; color: #7a1f2b; }
h3 { font-size: 12pt; line-height: 1.3; margin: 4pt 0 8pt; color: #7a1f2b; }
h1,h2,h3,.Subtitle { break-after: avoid; } p { margin: 0 0 8pt; orphans: 3; widows: 3; overflow-wrap: anywhere; }
.Subtitle { font-size: 14pt; color: #7a1f2b; margin-bottom: 10pt; }
.Meta,.Caption { font-size: 9pt; color: #687078; margin-bottom: 5pt; }
.Meta:has(+ h3) { margin-top: 15pt; break-after: avoid; }
a { color: #7a1f2b; text-decoration: underline; }
table { width: 100%; border-collapse: collapse; margin: 10pt 0 16pt; font-size: 9pt; table-layout: fixed; }
thead { display: table-header-group; } tr { break-inside: avoid; }
th { background: #7a1f2b; color: white; text-align: left; } td,th { padding: 8pt; vertical-align: top; border-bottom: 1px solid #e8dde0; overflow-wrap: anywhere; }
td p,th p { margin: 0; } tbody tr:nth-child(odd) { background: #f7f3f4; }
.summary { margin: 14pt 0; } .summary td { padding: 10pt; border: 0; }
.summary h2 { font-size: 24pt; border: 0; padding: 0; margin: 0 0 4pt; }
.summary .Meta { font-size: 8.5pt; line-height: 1.3; }
</style></head><body>''' + ''.join(parts) + '</body></html>'


def render_pdf(doc):
    from playwright.sync_api import sync_playwright
    from .drive_store import StorageError
    try:
        with sync_playwright() as playwright:
            executable = shutil.which('chromium') or shutil.which('chromium-browser')
            if not executable and os.name == 'nt' and not Path(playwright.chromium.executable_path).exists():
                edge = Path(os.environ.get('PROGRAMFILES(X86)', 'C:/Program Files (x86)')) / 'Microsoft/Edge/Application/msedge.exe'
                if edge.exists():
                    executable = str(edge)
            options = {'executable_path': executable} if executable else {}
            with playwright.chromium.launch(headless=True, **options) as browser:
                page = browser.new_page(java_script_enabled=False)
                page.route('**/*', lambda route: route.abort())
                page.set_content(document_html(doc), wait_until='load')
                return page.pdf(format='A4', print_background=True, display_header_footer=True,
                    margin={'top':'22mm', 'bottom':'20mm', 'left':'21mm', 'right':'21mm'},
                    header_template='<div style="font:8px Arial;color:#687078;width:100%;margin:0 21mm">СКБ ИНДУКЦИЯ / КОНКУРЕНТНАЯ АНАЛИТИКА</div>',
                    footer_template='<div style="font:8px Arial;color:#687078;width:100%;text-align:right;margin:0 21mm">Аналитическая записка · <span class="pageNumber"></span></div>')
    except Exception:
        raise StorageError('Не удалось создать PDF. Проверьте установку Chromium на сервере и повторите выпуск.') from None

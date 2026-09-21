"""Render the same frozen payload as the preview; archive proofs are inert."""
from copy import deepcopy
import hashlib
from html import escape
from io import BytesIO
import json
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from .document_style import EditorialStyle, render_pdf
from src.utils import parse_report_period
from .coverage import monitored

KINDS = {'news':'Новости сайтов','telegram':'Новости Telegram-каналов','products':'Продукция и предложения'}
STATES = {'success':'Проверен','partial':'Сбор неполный','error':'Ошибка'}


def render(snapshot, evidence, competitors, *, include_pdf=False):
    """evidence maps SHA-256 to bytes already obtained by an authorized worker."""
    settings=SimpleNamespace(report_primary_color='7A1F2B',report_muted_color='666666',report_font_name='Arial',report_table_fill='F3E6E9')
    builder=EditorialStyle(None,settings)
    doc=Document()
    builder._setup_document(doc)
    builder._title(doc,[SimpleNamespace(name=c['name']) for c in competitors.values()],parse_report_period(snapshot['period']))
    doc.add_paragraph(f"Сохранённая редакция: {snapshot['revision']}", style='Meta')
    if snapshot['incomplete']:
        doc.add_paragraph('СБОР НЕПОЛНЫЙ. Отсутствие материалов по непроверенным источникам не означает отсутствие новостей.').runs[0].bold=True
    builder.summary(doc,snapshot)
    doc.add_paragraph('Заголовки и описания отредактированы аналитиком. Даты, ссылки и сведения о проверке сохранены из первоисточников.')
    files={}
    portable=deepcopy(snapshot)
    for item in portable['items']:
        item['source']['evidence_files']=[]
        for asset_id in item['source']['evidence_ids']:
            content=evidence[asset_id]
            if hashlib.sha256(content).hexdigest()!=asset_id:
                raise ValueError('Контрольная сумма доказательства не совпадает.')
            # Preserve exact original bytes in a non-executable file. The linked
            # view displays escaped HTML with CSP and no scripts or external loads.
            raw='evidence/'+asset_id+'.html.txt'
            view='evidence/'+asset_id+'.html'
            files[raw]=content
            files[view]=('<!doctype html><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" '
                'content="default-src \'none\'; script-src \'none\'; sandbox"><title>Сохранённый первоисточник</title><pre>'
                +escape(content.decode('utf-8',errors='replace'))+'</pre>').encode()
            item['source']['evidence_files'].append({'original':raw,'view':view,'sha256':asset_id})
    for number,(kind,label) in enumerate(KINDS.items(),1):
        builder._heading(doc,f'{number}. {label}')
        selected=[i for i in portable['items'] if i['source']['kind']==kind]
        if not selected:
            doc.add_paragraph('Материалы не отобраны.')
            continue
        for item in selected:
            builder.article(doc,item)
    builder._heading(doc,'Выводы аналитика')
    for paragraph in (snapshot['conclusions'] or 'Выводы не добавлены.').splitlines():
        doc.add_paragraph(paragraph)
    builder._heading(doc,'Полнота проверки источников')
    rows=[['Конкурент','Источник','Состояние проверки']]
    for check in monitored(snapshot['checks']):
        status=STATES.get(check['status'],check['status'])
        if check['status']=='success' and check.get('items')==0:
            status+='; публикаций нет'
        rows.append([check.get('competitor_name',check['competitor_code']), check['url'],status])
    builder._table(doc,rows)
    builder._heading(doc,'Сохранённые доказательства')
    doc.add_paragraph('Относительные ссылки открываются после распаковки ZIP. Сохранённые страницы показаны как текст без выполнения скриптов.')
    for item in portable['items']:
        paragraph=doc.add_paragraph(item['title']+' — ')
        for index,proof in enumerate(item['source']['evidence_files'],1):
            builder._add_hyperlink(paragraph,proof['view'],f'Снимок {index} ')
    builder._heading(doc,'Подготовка и согласование')
    builder._signature_block(doc)
    doc.add_paragraph(f"Идентификатор черновика: {snapshot['draft_id']} · редакция {snapshot['revision']}", style='Meta')
    doc.core_properties.title='Аналитическая записка по конкурентам'
    doc.core_properties.author='СКБ «Индукция»'
    output=BytesIO(); doc.save(output)
    word=output.getvalue()
    files['report.docx']=word
    if include_pdf:
        files['report.pdf']=render_pdf(doc)
    files['manifest.json']=json.dumps(portable,ensure_ascii=False,indent=2).encode()
    bundle=BytesIO()
    with ZipFile(bundle,'w',ZIP_DEFLATED) as archive:
        for name,content in sorted(files.items()):
            archive.writestr(name,content)
    return word,bundle.getvalue()


def render_files(snapshot, evidence, competitors):
    word, bundle = render(snapshot, evidence, competitors, include_pdf=True)
    with ZipFile(BytesIO(bundle)) as archive:
        pdf = archive.read('report.pdf')
    return {'docx_id': word, 'pdf_id': pdf, 'bundle_id': bundle}

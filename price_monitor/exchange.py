from __future__ import annotations

import csv
from dataclasses import asdict
from io import BytesIO, StringIO
import re
import zipfile

from openpyxl import Workbook, load_workbook

from .models import Rule, normalize
from .sources import SOURCES, validate_url

COLUMNS = ["source", "manufacturer", "article", "product_url", "url_template"]
ALIASES = {"сайт": "source", "источник": "source", "производитель": "manufacturer", "артикул": "article", "модель": "article", "ссылка": "product_url", "шаблон": "url_template"}
MAX_ROWS = 10000
MAX_BYTES = 10 * 1024 * 1024


def read_table(data: bytes, filename: str) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ValueError("Максимальный размер файла — 10 МБ")
    if filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                    raise ValueError("Распакованный Excel превышает 80 МБ")
            wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = wb["Задания"] if "Задания" in wb.sheetnames else wb.worksheets[0]
                rows = []
                for cells in sheet.iter_rows():
                    if any(c.data_type == "f" for c in cells):
                        raise ValueError("Формулы в заданиях запрещены. Вставьте значения как текст")
                    if len(cells) > 20:
                        raise ValueError("В файле слишком много столбцов (максимум 20)")
                    rows.append([c.value for c in cells])
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError(f"Не более {MAX_ROWS} заданий за запуск")
            finally:
                wb.close()
        except (zipfile.BadZipFile, KeyError) as e:
            raise ValueError("Не удалось прочитать XLSX") from e
    elif filename.lower().endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1251")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(StringIO(text), dialect))
        if len(rows) > MAX_ROWS + 1:
            raise ValueError(f"Не более {MAX_ROWS} заданий за запуск")
    else:
        raise ValueError("Поддерживаются только XLSX и CSV")
    if not rows:
        raise ValueError("Файл пуст")
    headers = [ALIASES.get(str(c or "").strip().lower(), str(c or "").strip().lower()) for c in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if len(set(headers)) != len(headers) or any(not h or h not in COLUMNS for h in headers):
        raise ValueError("Проверьте заголовки: нужны уникальные столбцы из " + ", ".join(COLUMNS))
    if not {"source", "manufacturer", "article"}.issubset(headers):
        raise ValueError("Обязательные столбцы: source, manufacturer, article")
    if not {"product_url", "url_template"}.intersection(headers):
        raise ValueError("Нужен столбец product_url или url_template")
    result = []
    for n, row in enumerate(rows[1:], start=2):
        if not any(str(c or "").strip() for c in row):
            continue
        if any(str(c or "").strip() for c in row[len(headers):]):
            raise ValueError(f"Строка {n}: значения без заголовка")
        item = {k: str(v).strip() if v is not None else "" for k, v in zip(headers, row)}
        item["_row"] = n
        result.append(item)
    return result


def validate_rows(rows: list[dict]) -> tuple[list[Rule], list[dict]]:
    rules, errors, seen = [], [], set()
    if not rows:
        return [], [{"Строка": 0, "Ошибка": "Нет заданий"}]
    if len(rows) > MAX_ROWS:
        return [], [{"Строка": 0, "Ошибка": f"Не более {MAX_ROWS} заданий"}]
    for i, row in enumerate(rows, start=2):
        try:
            values = {k: str(row.get(k, "") or "").strip() for k in COLUMNS}
            source = values["source"].lower()
            source = next((sid for sid,s in SOURCES.items() if source in {sid, s.host, s.label.lower()}), source)
            if source not in SOURCES:
                raise ValueError("Неизвестный источник: " + values["source"])
            values["source"] = source
            brand = values["manufacturer"]
            brand_key = re.sub(r"[^\w]", "", normalize(brand))
            brand_key = {"IFMELECTRONIC": "IFM", "PEPPERLFUCHS": "PEPPERLFUCHS"}.get(brand_key, brand_key)
            canonical = next((b for b in SOURCES[source].brands if re.sub(r"[^\w]", "", normalize(b)) == brand_key), None)
            if not canonical:
                raise ValueError("Производитель для источника: " + ", ".join(SOURCES[source].brands))
            values["manufacturer"] = canonical
            article = values["article"]
            if not article or len(article) > 150 or any(ord(c) < 32 for c in article) or article.startswith(("=", "+", "@")):
                raise ValueError("Укажите артикул текстом (до 150 символов, без формул)")
            if bool(values["product_url"]) == bool(values["url_template"]):
                raise ValueError("Заполните ровно одно: product_url или url_template")
            template = values["url_template"]
            if template and (template.count("{article}") != 1 or re.search(r"[{}]", template.replace("{article}", ""))):
                raise ValueError("В шаблоне нужен ровно один параметр {article}")
            rule = Rule(**values)
            if len(rule.url) > 2048:
                raise ValueError("Ссылка длиннее 2048 символов")
            validate_url(source, rule.url)
            if rule.key in seen:
                raise ValueError("Повторное задание на ту же модель и ссылку")
            seen.add(rule.key)
            rules.append(rule)
        except ValueError as e:
            errors.append({"Строка": row.get("_row", i), "Ошибка": str(e)})
    return rules, errors


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def csv_bytes(rows: list[dict], columns: list[str] | None = None) -> bytes:
    columns = columns or (list(rows[0]) if rows else COLUMNS)
    out = StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=columns, delimiter=";", extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe_cell(v) for k, v in row.items()})
    return out.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows: list[dict], columns: list[str] | None = None, sheet_name="Результаты") -> bytes:
    columns = columns or (list(rows[0]) if rows else COLUMNS)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(columns)
    for row in rows:
        ws.append([safe_cell(row.get(c)) for c in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(65, max(18, len(str(col[0].value or "")) + 3))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def rules_rows(rules):
    return [asdict(r) for r in rules]

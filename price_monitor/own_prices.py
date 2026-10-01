"""Versioned SKB prices: exact article identity, original net amount and provenance."""
from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook

from .matching_normalize import model_key
from .models import utcnow

SCHEMA = """
CREATE TABLE IF NOT EXISTS own_price_imports (
 batch_id TEXT PRIMARY KEY, source_name TEXT NOT NULL, source_hash TEXT NOT NULL,
 effective_date TEXT NOT NULL, imported_at TEXT NOT NULL, vat_rate TEXT NOT NULL,
 row_count INTEGER NOT NULL, matched_count INTEGER NOT NULL, report_json TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS own_product_prices (
 batch_id TEXT NOT NULL REFERENCES own_price_imports(batch_id), article TEXT NOT NULL,
 source_model TEXT NOT NULL, catalog_id TEXT NOT NULL DEFAULT '', catalog_model TEXT NOT NULL DEFAULT '',
 net_price TEXT NOT NULL, gross_price TEXT NOT NULL, currency TEXT NOT NULL,
 unit TEXT NOT NULL, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
 PRIMARY KEY(batch_id,article)
);
CREATE INDEX IF NOT EXISTS own_prices_catalog ON own_product_prices(catalog_id);
"""


def read_prices(data: bytes, source_name: str, effective_date: str, matcher, vat_rate='22'):
    """Validate the entire workbook before any write. No approximate product links."""
    effective_date = date.fromisoformat(effective_date).isoformat()
    rate = Decimal(str(vat_rate))
    if not rate.is_finite() or not 0 <= rate <= 100:
        raise ValueError('Ставка НДС должна быть от 0 до 100%.')
    source_hash = sha256(data).hexdigest()
    batch_id = sha256(f'{source_hash}|{effective_date}|{rate}'.encode()).hexdigest()
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    rows, seen, errors = [], set(), []
    try:
        for sheet in workbook:
            columns = None
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if not any(v is not None for v in values):
                    continue
                labels = [str(v or '').strip().casefold() for v in values]
                if columns is None:
                    if all(k in labels for k in ('наименование', 'артикул', 'цена')):
                        columns = [labels.index(k) for k in ('наименование', 'артикул', 'цена')]
                    continue
                name, article, raw_price = (values[i] for i in columns)
                article = str(article).strip() if article is not None else ''
                if not name or not article:
                    errors.append(f'{sheet.title}, строка {number}: нет наименования или артикула')
                    continue
                if article in seen:
                    errors.append(f'{sheet.title}, строка {number}: повтор артикула {article}')
                    continue
                try:
                    price = Decimal(str(raw_price).replace('\xa0', '').replace(' ', '').replace(',', '.'))
                    if not price.is_finite() or price <= 0:
                        raise InvalidOperation
                except (InvalidOperation, ValueError):
                    errors.append(f'{sheet.title}, строка {number}: некорректная цена')
                    continue
                seen.add(article)
                # The 1C article is authoritative, including leading zeroes.
                sensor = matcher.resolve(article)
                if sensor is None:
                    candidate = matcher.resolve(str(name))
                    sensor = candidate if candidate and model_key(candidate.article) == model_key(article) else None
                rows.append(dict(batch_id=batch_id, article=article, source_model=str(name),
                    catalog_id=sensor.id if sensor else '', catalog_model=sensor.model if sensor else '',
                    net_price=format(price, 'f'),
                    gross_price=format((price * (1 + rate / 100)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP), 'f'),
                    currency='RUB', unit='piece' if sensor else '', source_sheet=sheet.title, source_row=number))
    finally:
        workbook.close()
    if errors or not rows:
        raise ValueError('; '.join(errors[:12]) if errors else 'Не найден лист с колонками Наименование, Артикул, Цена.')
    matched = sum(bool(r['catalog_id']) for r in rows)
    report = dict(rows=len(rows), matched=matched, unmatched=len(rows)-matched,
                  net_sum=str(sum(Decimal(r['net_price']) for r in rows)),
                  gross_sum=str(sum(Decimal(r['gross_price']) for r in rows)),
                  tax_basis='net', vat_rate=str(rate), effective_date=effective_date)
    metadata = dict(batch_id=batch_id, source_name=Path(source_name).name, source_hash=source_hash,
                    effective_date=effective_date, imported_at=utcnow(), vat_rate=str(rate),
                    row_count=len(rows), matched_count=matched, report_json=json.dumps(report, ensure_ascii=False))
    return metadata, rows


class OwnPrices:
    def __init__(self, repository):
        self.repo = repository

    def initialize(self):
        self.repo.batch([sql for sql in SCHEMA.split(';') if sql.strip()])

    def import_rows(self, metadata, rows):
        """One atomic import; rerunning the same file/date/rate is idempotent."""
        if len(rows) != metadata['row_count'] or any(r['batch_id'] != metadata['batch_id'] for r in rows):
            raise ValueError('Состав строк не соответствует проверенному импорту.')
        if self.repo.batch('SELECT batch_id FROM own_price_imports WHERE batch_id=%(batch)s', {'batch':metadata['batch_id']}):
            return False
        params = dict(metadata)
        sql = ['''INSERT INTO own_price_imports
            (batch_id,source_name,source_hash,effective_date,imported_at,vat_rate,row_count,matched_count,report_json)
            VALUES(%(batch_id)s,%(source_name)s,%(source_hash)s,%(effective_date)s,%(imported_at)s,%(vat_rate)s,%(row_count)s,%(matched_count)s,%(report_json)s)''']
        columns = ('batch_id','article','source_model','catalog_id','catalog_model','net_price','gross_price','currency','unit','source_sheet','source_row')
        for n, row in enumerate(rows):
            params.update({f'{key}_{n}':row[key] for key in columns})
            sql.append('INSERT INTO own_product_prices ('+','.join(columns)+') VALUES ('+
                       ','.join(f'%({key}_{n})s' for key in columns)+')')
        self.repo.batch(sql, params)
        return True

    def current(self, as_of=None):
        rows = self.repo.batch('''SELECT p.*,i.effective_date,i.vat_rate,i.source_name,i.source_hash,i.imported_at
            FROM own_product_prices p JOIN own_price_imports i ON i.batch_id=p.batch_id
            WHERE i.effective_date<=%(today)s ORDER BY i.effective_date,i.imported_at,p.article''',
            {'today':as_of or date.today().isoformat()})
        latest = {}
        for row in rows:
            latest[row['article']] = row
        return list(latest.values())

    def catalog_prices(self, as_of=None):
        return {row['catalog_id']:row for row in self.current(as_of) if row['catalog_id']}

    def imports(self):
        return self.repo.batch('SELECT * FROM own_price_imports ORDER BY effective_date DESC,imported_at DESC')

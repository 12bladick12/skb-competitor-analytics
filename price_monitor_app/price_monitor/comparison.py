from __future__ import annotations

from datetime import datetime
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from .exchange import read_table, ALIASES
from .library import price_value
from .models import normalize
from .sources import SOURCES,validate_url

COLUMNS=['source','manufacturer','article','our_price','our_currency','our_article','product_url','note']
LABELS={'source':'Источник','manufacturer':'Производитель','article':'Артикул конкурента',
        'our_price':'Наша цена','our_currency':'Валюта нашей цены','our_article':'Наш артикул','product_url':'Ссылка на товар','note':'Примечание'}
ALIASES={**ALIASES,**{v.casefold():k for k,v in LABELS.items()},'валюта':'our_currency'}


def template_bytes():
    book=Workbook();sheet=book.active;sheet.title='Сравнение'
    sheet.append(COLUMNS);sheet.append(['']*len(COLUMNS));sheet.freeze_panes='A2'
    sheet.auto_filter.ref='A1:H1001'
    for column in ('A','B','C','F','G','H'):
        for row in range(2,1002):sheet[f'{column}{row}'].number_format='@'
    for row in range(2,1002):sheet[f'D{row}'].number_format='#,##0.00'
    for col,width in zip('ABCDEFGH',(18,22,36,18,20,32,65,50)):
        sheet.column_dimensions[col].width=width
    for formula,column in [('"'+','.join(SOURCES)+'"','A'),('"RUB,USD,EUR,CNY"','E')]:
        validation=DataValidation(type='list',formula1=formula,allow_blank=True)
        validation.errorTitle='Выберите значение из списка';validation.error='Используйте значение справочника'
        validation.showErrorMessage=True;sheet.add_data_validation(validation);validation.add(f'{column}2:{column}1001')
    guide=book.create_sheet('Инструкция');guide.append(['Столбец','Как заполнять'])
    notes={
        'source':'Обязательно. Код источника из листа «Источники», например sensoren.',
        'manufacturer':'Обязательно. Производитель из справочника, например ifm.',
        'article':'Обязательно. Точный артикул конкурента. Сохраняйте ведущие нули, дефисы и скобки.',
        'our_price':'Необязательно. Наша текущая цена числом больше 0, без формул. Пусто = цена не задана.',
        'our_currency':'RUB, USD, EUR или CNY. Пусто = RUB. Сравнение с нашей ценой только в той же валюте.',
        'our_article':'Необязательно. Наш артикул; связь задаёте вы вручную.',
        'product_url':'Необязательно. HTTPS-ссылка из базы. Обязательна, если у артикула несколько карточек.',
        'note':'Необязательно. Условия цены, НДС, количество, пояснение соответствия.'}
    for key in COLUMNS:guide.append([key,notes[key]])
    for note in [
        'Заполняйте лист «Сравнение». Лист «Пример» не импортируется.',
        'Сначала соберите товары в базу приложения. Импорт ищет точное совпадение, аналоги не подбирает.',
        'При ошибке хотя бы одной строки сохранение всего файла блокируется. Исправьте файл и загрузите снова.',
        'Загрузка обновляет нашу цену и примечание для перечисленных позиций. Остальной список сохраняется.',
        'Наша цена — текущий ориентир, а не историческая серия. НДС, упаковка и курсы валют автоматически не пересчитываются.',
        'История и список сравнения общие для всех посетителей приложения.'
    ]:guide.append(['Правило',note])
    sources=book.create_sheet('Источники');sources.append(['source','manufacturer','Сайт'])
    for source,spec in SOURCES.items():
        for brand in spec.brands:sources.append([source,brand,'https://'+spec.host])
    example=book.create_sheet('Пример');example.append(COLUMNS)
    example.append(['sensoren','ifm','SI5000','','RUB','','','Введите свою цену; модель должна быть в базе'])
    for ws in book:
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='7A1F2B');cell.font=Font(color='FFFFFF',bold=True)
        ws.freeze_panes='A2';ws.row_dimensions[1].height=28
        if ws!=sheet:
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=90 if ws==guide and col[0].column==2 else 26
            for row in ws.iter_rows(min_row=2):
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
                ws.row_dimensions[row[0].row].height=44 if ws==guide else 30
    output=BytesIO();book.save(output);return output.getvalue()


def read_comparisons(data,name):
    rows=read_table(data,name,columns=COLUMNS,aliases=ALIASES,require_url=False,sheet_name='Сравнение')
    output=[];errors=[]
    for row in rows:
        try:
            source=str(row['source']).strip().casefold()
            source=next((key for key,spec in SOURCES.items() if source in (key,spec.label.casefold(),spec.host)),source)
            if source not in SOURCES:raise ValueError('Неизвестный источник')
            brandkey=re.sub(r'[^\w]','',normalize(row['manufacturer'])).replace('IFMELECTRONIC','IFM')
            brand=next((x for x in SOURCES[source].brands if re.sub(r'[^\w]','',normalize(x))==brandkey),None)
            if not brand:raise ValueError('Производитель не относится к выбранному источнику')
            article=str(row['article']).strip()
            if not article or len(article)>240 or article.startswith(('=','@','+')):raise ValueError('Укажите точный артикул')
            currency=(row.get('our_currency') or 'RUB').upper()
            if currency not in ('RUB','USD','EUR','CNY'):raise ValueError('Валюта: RUB, USD, EUR, CNY')
            url=row.get('product_url','')
            if url:validate_url(source,url)
            output.append({**row,'source':source,'manufacturer':brand,'article':article,
                'our_price':price_value(row.get('our_price')),'our_currency':currency})
        except (ValueError,KeyError) as exc:errors.append({'Строка':row['_row'],'Ошибка':str(exc)})
    if not rows:errors.append({'Строка':2,'Ошибка':'Заполните лист «Сравнение»'})
    return output,errors


def series(history,currency):
    points=[];segment=0
    for item in history:
        if item['status']!='priced' or item.get('currency')!=currency or item.get('price') is None:
            segment+=1;continue
        points.append({'date':datetime.fromisoformat(item['checked_at']),'price':float(item['price']),
                       'segment':str(segment),'run':item['run_id'],'checked_at':item['checked_at']})
    return points


def metrics(points,our_price,our_currency,currency):
    latest=points[-1]['price'] if points else None
    change=(latest/points[0]['price']-1)*100 if len(points)>1 and points[0]['price'] else None
    reference=float(our_price) if our_price and our_currency==currency else None
    gap=(latest/reference-1)*100 if latest is not None and reference else None
    return latest,change,reference,gap

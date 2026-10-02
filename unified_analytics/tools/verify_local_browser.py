"""Read-only browser smoke check of the running local app."""
import argparse
import json
from pathlib import Path
import re
import sqlite3
import time

from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--url', default='http://127.0.0.1:8530/')
    parser.add_argument('--require-data', action='store_true')
    args = parser.parse_args()
    output = ROOT / 'data/local/validation'
    output.mkdir(parents=True, exist_ok=True)
    report = []
    with sync_playwright() as api:
        browser = api.chromium.launch(headless=True)
        page = browser.new_page(viewport={'width': 1440, 'height': 1000}, device_scale_factor=1)
        chart_warnings = []
        page.on('console', lambda message: chart_warnings.append(message.text) if 'Infinite extent' in message.text else None)

        def check(section):
            started = time.monotonic()
            page.goto(args.url + '?workspace=prices&price_section=' + section, wait_until='domcontentloaded')
            expect(page.locator('[data-testid="stApp"]')).to_have_attribute(
                'data-test-script-state', 'notRunning', timeout=120000)
            expect(page.get_by_text('Локальный режим · цены сохраняются на этом компьютере', exact=True)).to_be_visible()
            assert page.locator('[data-testid="stException"]').count() == 0
            alerts = page.locator('[data-testid="stAlertContainer"]')
            assert 'Не удалось открыть данные' not in ' '.join(alerts.all_text_contents())
            assert 'Не удалось выполнить действие' not in ' '.join(alerts.all_text_contents())
            assert not re.search(r'\?{3,}', page.locator('body').inner_text())
            if section == 'changes':
                expect(page.get_by_role('heading', name='Изменения цен', level=1)).to_be_visible()
            if args.require_data and section == 'products':
                metric = page.locator('[data-testid="stMetric"]').filter(has_text='Товаров в базе')
                value = metric.locator('[data-testid="stMetricValue"]').inner_text()
                assert int(''.join(c for c in value if c.isdigit())) > 0
                expect(page.locator('[data-testid="stDataFrame"]').first).to_be_visible()
            if args.require_data and section == 'runs':
                expect(page.get_by_role('combobox', name='Запуск', exact=True)).to_be_visible()
            page.screenshot(path=str(output / (section + '.png')), full_page=False)
            report.append({'section': section, 'seconds': round(time.monotonic() - started, 2), 'exceptions': 0})

        for section in ('home', 'collect', 'products', 'changes', 'runs', 'compare'):
            check(section)
        search = page.get_by_role('combobox', name='Поиск по номенклатуре или артикулу', exact=True)
        expect(search).to_be_visible()
        search.fill('И27')
        expect(page.get_by_role('option').first).to_be_visible(timeout=30000)
        option = page.get_by_role('option').filter(has_text='СКБ Индукция').first
        expect(option).to_be_visible(timeout=30000)
        selected = option.inner_text()
        option.click()
        search.press('Escape')
        expect(page.get_by_role('combobox', name='Сравнивать цену с', exact=True).first).to_be_visible(timeout=120000)
        expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state', 'notRunning', timeout=120000)
        assert page.locator('[data-testid="stException"]').count() == 0
        differences = page.locator('.comparison-delta').all_text_contents()
        assert not any(re.search(r'[-−+]\s*\d', text) for text in differences), differences
        assert not any('Выключатель индуктивный' in text for text in page.locator('.comparison-model').all_text_contents())
        assert not chart_warnings, chart_warnings
        report.append({'search': 'И27', 'selected': selected})
        page.screenshot(path=str(output / 'comparison-selected.png'))
        first_chart = page.locator('[data-testid="stVegaLiteChart"]').first
        if first_chart.count():
            first_chart.scroll_into_view_if_needed()
            page.screenshot(path=str(output / 'price-chart.png'))
        report.append({'comparison_unsigned_differences': len(differences), 'chart_warnings': len(chart_warnings)})
        collapse = page.locator('[data-testid="stSidebarCollapseButton"] button')
        if collapse.is_visible():
            collapse.click()
        page.set_viewport_size({'width': 430, 'height': 900})
        page.wait_for_timeout(600)
        page.screenshot(path=str(output / 'comparison-mobile.png'))
        assert not page.evaluate('document.documentElement.scrollWidth > window.innerWidth')
        report.append({'mobile_width': 430, 'page_overflow': False})
        if args.require_data:
            page.set_viewport_size({'width': 1440, 'height': 1000})
            page.goto(args.url + '?workspace=prices&price_section=products', wait_until='domcontentloaded')
            query = page.get_by_role('textbox', name='Поиск номенклатуры', exact=True)
            query.fill('SI5000')
            query.press('Enter')
            page.get_by_text('Выгрузить найденные товары', exact=True).click()
            page.get_by_role('button', name='Подготовить выгрузку', exact=True).click()
            download_button = page.get_by_role('button', name='Скачать XLSX', exact=True)
            expect(download_button).to_be_visible(timeout=60000)
            with page.expect_download() as download:
                download_button.click()
            destination = output / 'products-export.xlsx'
            download.value.save_as(destination)
            from openpyxl import load_workbook
            book = load_workbook(destination, read_only=True)
            try:
                assert book.active.max_row > 1
                assert any('SI5000' in str(value) for row in book.active.iter_rows(values_only=True) for value in row)
                report.append({'export': 'XLSX', 'sheets': book.sheetnames, 'rows': book.active.max_row - 1})
            finally:
                book.close()
            with sqlite3.connect((ROOT / 'data/local/prices.sqlite3').as_uri() + '?mode=ro', uri=True) as connection:
                model = connection.execute("SELECT id FROM rules WHERE source='teko' AND article LIKE 'Выключатель индуктивный %' LIMIT 1").fetchone()
            if model:
                page.goto(args.url + '?workspace=prices&price_section=history&model=' + str(model[0]), wait_until='domcontentloaded')
                expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state', 'notRunning', timeout=60000)
                assert not page.locator('[data-testid="stException"]').count()
                assert 'Выключатель индуктивный' not in ' '.join(page.locator('h3').all_text_contents())
                assert not chart_warnings, chart_warnings
                page.screenshot(path=str(output / 'teko-history.png'))
                report.append({'teko_history': model[0], 'chart_warnings': 0})
        browser.close()
    (output / 'browser_checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()

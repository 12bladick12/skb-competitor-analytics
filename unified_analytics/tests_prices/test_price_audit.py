import csv,json,unittest,tempfile
from io import BytesIO,StringIO
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from price_monitor.details import manufacturer,attributes,parse_catalog_product
from price_monitor.catalog import discover,seeds
from price_monitor.exchange import csv_bytes,xlsx_bytes
from price_monitor.product_exports import export_tables
from price_monitor.transport import SourceClient,FetchError,challenge
from price_monitor.storage import Store
from price_monitor.library import Library
from price_monitor.catalog_schema import document
from price_monitor.scope import refresh_scope
from price_monitor.adapters import ADAPTERS
from price_monitor.models import Rule

class AuditChecks(unittest.TestCase):
    def test_teko_brand_is_explicit(self):
        def page(brand):return BeautifulSoup('<h1>ТЕКО sample</h1><div class="product-item-detail-properties"><div class="one_prop"><div class="name">Бренд</div><div class="value">'+brand+'</div></div></div>','html.parser')
        self.assertEqual(manufacturer('teko',page('ТЕКО')),'ТЕКО')
        self.assertEqual(manufacturer('teko',page('Autonics')),'Autonics')
        self.assertEqual(manufacturer('teko',BeautifulSoup('<footer>НПК ТЕКО</footer>','html.parser')),'')

    def test_teko_catalog_and_manual_jobs_reject_other_brands(self):
        own='''<h1>Соединитель ТЕКО CP S254R-3 (PG9)</h1><div class="catalog-detail">
            <div class="product-item-detail-properties"><div class="one_prop"><div class="name">Бренд</div><div class="value">ТЕКО</div></div></div>
            <div class="catalog-detail-price"><div>1708 руб.</div></div></div>'''
        url='https://teko-com.ru/catalog/product/cp-s254r-3-pg9/'
        results,reason=parse_catalog_product('teko',own,url,['ТЕКО'])
        self.assertEqual(reason,'');self.assertEqual(results[0][1].status,'priced')
        foreign=own.replace('<div class="value">ТЕКО','<div class="value">Autonics')
        self.assertEqual(parse_catalog_product('teko',foreign,url,['ТЕКО']),([], 'outside_scope'))
        result=ADAPTERS['teko'].parse(Rule('teko','ТЕКО','CP S254R-3 (PG9)',url),foreign,url)
        self.assertEqual(result.status,'identity_mismatch');self.assertIsNone(result.price)

    def test_sensoren_unselected_manufacturer_is_excluded(self):
        html='<h1>Balluff ABC</h1><div class="product-info"><div class="product-info__brand-name">Balluff</div></div>'
        self.assertEqual(parse_catalog_product('sensoren',html,'https://sensoren.ru/product/abc/',['ifm']),([], 'outside_scope'))

    def test_sensoren_challenge_stops(self):
        body='<script>document.cookie="RCPC=example; path=/";document.location.href="/?attempt=1";</script>'
        self.assertTrue(challenge(body))
        c=SourceClient('sensoren')
        try:
            with self.assertRaises(FetchError) as cm:c._check_status(503,{},body,'https://sensoren.ru/catalog/')
            self.assertTrue(cm.exception.stop_source)
            self.assertIn('RCPC',str(cm.exception))
        finally:c.close()

    def test_brand_navigation_scope(self):
        self.assertEqual(seeds('sensoren',['ifm'])[0],('listing','https://sensoren.ru/brands/ifm_electronic/'))
        links=discover('sensoren','listing','''<a href="/brands/ifm_electronic/?PAGEN_1=2">next</a>
            <a href="/brands/balluff/">other</a><a href="/catalog/all/">menu</a>
            <a href="/catalog/sensors/brand_ifm_electronic/">own category</a>
            <a href="/catalog/sensors/brand_balluff/">other category</a>
            <a href="/product/datchik_ifm_abc/">IFM ABC</a>
            <a href="/product/datchik_balluff_def/">Balluff DEF</a>''','https://sensoren.ru/brands/ifm_electronic/',['ifm'])
        self.assertEqual(len(links),3)

    def test_nested_sensoren_properties(self):
        soup=BeautifulSoup('<ul class="characteristics-all"><li><b>Напряжение:</b><span>24 В</span></li></ul>','html.parser')
        self.assertEqual(attributes('sensoren',soup)[0]['name'],'Напряжение')

    def test_exports_include_all_rows_and_sheet(self):
        rows=[{'article':'A','_specifications':{'attributes':[{'name':'Материал','value':'Сталь','group':''}]}},
              {'article':'B','_specifications':{'attributes':[{'name':'Выход','value':'=1+1','group':''}]}}]
        flat,props=export_tables(rows)
        parsed=list(csv.DictReader(StringIO(csv_bytes(flat).decode('utf-8-sig')),delimiter=';'))
        self.assertIn('Характеристика: Выход',parsed[0])
        self.assertEqual(parsed[1]['Характеристика: Выход'],"'=1+1")
        wb=load_workbook(BytesIO(xlsx_bytes(flat,extra_sheets={'Характеристики':props})))
        self.assertEqual(wb['Характеристики'].max_row,3)
        self.assertEqual(wb['Результаты'].max_row,3)

    def test_scope_backfill_and_snapshot_export(self):
        root=Path(__file__).resolve().parents[1]/'data'
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as tmp:
            store=Store(Path(tmp)/'prices.sqlite3')
            docs=[]
            with store.connect() as c:
                for rid,brand in [(1,'ТЕКО'),(2,'Autonics')]:
                    payload={'attributes':[{'name':'Бренд','value':brand,'group':''},{'name':'Материал','value':'Сталь','group':''}]}
                    fingerprint,encoded,_=document(json.dumps(payload))
                    docs.append(fingerprint)
                    c.execute("INSERT INTO rules VALUES(?,?,?,?,?,?,?,?)",(rid,str(rid),'teko','ТЕКО',str(rid),'https://teko-com.ru/catalog/product/'+str(rid)+'/','','2026-09-29'))
                    c.execute('INSERT INTO product_documents VALUES(?,?)',(fingerprint,encoded))
                    c.execute('INSERT INTO product_index VALUES(?,?,?,?,?,?,?)',(rid,'title','','',fingerprint,2,'2026-09-29'))
            self.assertEqual(refresh_scope(store.catalog),2)
            self.assertEqual(refresh_scope(store.catalog),0)
            library=Library(store.catalog)
            total,items=library.products(source='teko')
            self.assertEqual(total,1);self.assertEqual(items[0]['rule_id'],1)
            self.assertEqual(library.summary()['products'],1)
            rows=library.with_specifications([{'rule_id':1,'details_json':json.dumps({'ref':docs[0]})},{'rule_id':1,'details_json':'{}'}])
            self.assertEqual(len(rows[0]['_specifications']['attributes']),2)
            self.assertEqual(rows[1]['_specifications'],{})
            with store.connect() as c:c.execute("UPDATE product_index SET details_hash='changed' WHERE rule_id=1")
            self.assertEqual(library.products(source='teko')[0],0)

if __name__=='__main__':unittest.main()

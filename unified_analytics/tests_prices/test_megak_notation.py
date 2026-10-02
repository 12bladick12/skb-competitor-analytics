"""Boundary tests for published notation, source precedence and the real HTML tab."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from price_monitor.details import attributes
from price_monitor.megak_notation import decode, is_megak, SOURCE_URL
from price_monitor.matching_normalize import normalize_sensor
from price_monitor.matching import evaluate, matching_export
from price_monitor.exchange import xlsx_bytes
from tests_prices.test_inductive_matching import sensor


def product(model='PS2-18M68-8N11-K',props=None,**kw):
    return normalize_sensor({'manufacturer':'МЕГА-К','article':model,'props':props or {},**kw})


class DesignationChecks(unittest.TestCase):
    def test_standard_cable_model_and_published_defaults(self):
        p=product()
        self.assertEqual(p.family,'inductive')
        self.assertEqual({k:p.values[k] for k in ('diameter','length','sn','mount','output','function','voltage_type','vmin','vmax')},
                         dict(diameter=18.,length=68.,sn=8.,mount='non-flush',output='PNP',function='NO',voltage_type='DC',vmin=10.,vmax=30.))
        self.assertEqual((p.values['material'],p.values['ip'],p.values['tmin'],p.values['tmax'],p.values['load']),('brass',('67',),-25.,75.,250.))
        self.assertTrue(p.decoding['complete']);self.assertFalse(p.decoding['requires_review'])
        self.assertEqual(p.decoding['fields']['material']['application'],'filled')
        self.assertTrue(p.decoding['fields']['material']['default'])

    def test_manufacturer_example_and_special_application(self):
        p=product('PS2A-12M68-4B11-C4-T1')
        self.assertEqual((p.values['diameter'],p.values['sn'],p.values['tmin'],p.values['tmax']),(12.,4.,-40.,105.))
        self.assertEqual((p.values['wire_count'],p.values['pin_count'],p.values['connector']),(3.,4.,'m12'))
        self.assertTrue(p.decoding['requires_review'])

    def test_vb_new_format_and_cyrillic_lookalikes(self):
        p=product(' vb2 – 18МS68 – 8В31 – С4 – Т3Y05N1 ')
        self.assertTrue(p.decoding['complete'])
        self.assertEqual((p.values['material'],p.values['function'],p.values['tmin'],p.values['load']),('steel','NC',-45.,500.))
        self.assertEqual(p.profile,'general')

    def test_multidigit_circuit_and_last_digit_voltage(self):
        p=product('PS2-12M33-0,5B131-K')
        self.assertEqual((p.values['sn'],p.values['function'],p.values['output'],p.values['vmin']),(0.5,'NO/NC','PNP',10.))
        p=product('PS2-18M53-8N51-K')
        self.assertEqual((p.values['function'],p.values['wire_count']),('changeover',4.))

    def test_supply_codes_do_not_merge_ac_and_dc(self):
        p=product('PS2-18M68-8N72-K')
        self.assertEqual((p.values['output'],p.values['voltage_type'],p.values['vmin']),('2-wire','AC',35.))
        p=product('PS2-18M68-8N74-K')
        self.assertEqual(p.values['voltage_type'],'AC/DC');self.assertNotIn('vmin',p.values)
        self.assertTrue(p.decoding['requires_review'])
        for code in ('3','5'):
            p=product('PS2-18M68-8N1'+code+'-K')
            self.assertNotIn('vmin',p.values);self.assertNotIn('vmax',p.values)

    def test_special_case_number_is_not_diameter_length_or_brass(self):
        p=product('PS2-33-15N11-K')
        self.assertEqual(p.values['body_type'],'rectangular')
        for name in ('diameter','length','material'):self.assertNotIn(name,p.values)
        p=product('PS2-46-15N11-K')
        self.assertTrue(p.decoding['complete']);self.assertTrue(p.decoding['requires_review'])
        self.assertEqual(p.decoding['extras']['Номер корпуса'],'46')
        for name in ('body_type','diameter','length','material'):self.assertNotIn(name,p.values)

    def test_unknown_suffix_blocks_defaults_but_keeps_explicit_data(self):
        for model in ('PS2-18M68-8N11-K-T9','PS2-18M68-8N11-C99','PS2Q-18M68-8N11-K'):
            p=product(model)
            self.assertFalse(p.decoding['complete']);self.assertTrue(p.decoding['requires_review'])
            for name in ('material','tmin','tmax','ip','load'):self.assertNotIn(name,p.values)
            self.assertEqual(p.values['diameter'],18.)

    def test_truncated_unseparated_and_old_designations_are_not_guessed(self):
        for model in ('PS2-18M68-8N11-','PS2-18M68','ВБ2.18М.8.1.1.К','PS2-18M68-1511-K'):
            p=product(model)
            self.assertFalse(p.decoding['complete']);self.assertNotIn('load',p.values)

    def test_other_manufacturers_and_other_product_types_not_decoded(self):
        for row in ({'manufacturer':'ТЕКО','source':'megak'},{'manufacturer':'ifm'},{'source':'sensoren'},{}):
            p=normalize_sensor({'article':'PS2-18M68-8N11-K',**row})
            self.assertEqual(p.values,{})
            self.assertFalse(p.decoding.get('supported',False))
            self.assertFalse(p.decoding.get('fields',{}))
        self.assertTrue(is_megak({'source':'megak'}))
        self.assertTrue(is_megak({'manufacturer':'Mega-K','source':'sensoren'}))
        for prefix in ('PS1','PS3','VB5','PS9'):
            self.assertEqual(product(prefix+'-18M68-8N11-K').values,{})

    def test_cable_length_and_connector_are_independent(self):
        p=product('PS2-18M68-8N11-KPu02C4-T3')
        self.assertEqual(p.values['connection'],'cable+connector')
        self.assertEqual(p.decoding['extras']['Длина кабеля, м'],0.2)
        self.assertEqual(p.decoding['extras']['Материал кабеля'],'Полиуретан')
        self.assertNotIn('Длина кабеля, м',product('PS2-18M68-8N11-C4').decoding['extras'])
        self.assertNotIn('connection',product('PS2-18M68-8N11').values)

    def test_frequency_and_analog_execution_are_not_standard(self):
        p=product('PS2R1T4G-18M68-8N11-K')
        self.assertEqual(p.special,('speed',));self.assertNotIn('frequency',p.values)
        self.assertEqual(p.values['ip'],('68',))
        self.assertEqual(product('PS2D-18M68-8N91-K').special,('analog',))

    def test_thread_pitch_and_switching_frequency_are_never_inferred(self):
        p=product()
        self.assertNotIn('pitch',p.values);self.assertNotIn('frequency',p.values)
        self.assertEqual(evaluate(p,sensor(diameter=18.,sn=8.,mount='non-flush')).status,'review')

    def test_pipe_thread_and_smooth_housings_remain_distinct(self):
        self.assertEqual(product('PS2-18G68-8N11-K').values['body_type'],'pipe-threaded')
        self.assertEqual(product('PS2-18D68-8N11-K').values['body_type'],'smooth')


class EnrichmentChecks(unittest.TestCase):
    def test_explicit_card_wins_and_conflict_requires_review(self):
        p=product(props={'Расстояние срабатывания номинальное (Sn)':'10 мм','Типоразмер корпуса, мм':'M18x1','Частота переключения максимальная (f)':'500 Гц'})
        self.assertEqual(p.values['sn'],10.)
        self.assertEqual(p.decoding['fields']['sn']['application'],'conflict')
        self.assertEqual(evaluate(p,sensor(sn=10.,mount='non-flush',tmin=-25.,tmax=75.,load=250.)).status,'review')

    def test_card_overrides_omitted_defaults_without_false_conflict(self):
        p=product(props={'Материал корпуса':'Нержавеющая сталь','Ток нагрузки максимальный (Ie)':'0,3 А','Рабочая температура':'-40...+85 °C'})
        self.assertEqual((p.values['material'],p.values['load'],p.values['tmin']),('stainless',300.,-40.))
        self.assertFalse(p.decoding['requires_review'])
        self.assertEqual(p.decoding['fields']['load']['application'],'card_priority')

    def test_two_card_values_cannot_be_resolved_by_model_code(self):
        p=product(props={'Типоразмер корпуса, мм':'M18x1','Размер резьбы корпуса':'M12x1'})
        self.assertIn('diameter',p.conflicts);self.assertNotIn('diameter',p.values)
        self.assertEqual(p.decoding['fields']['diameter']['application'],'card_conflict')

    def test_placeholders_are_missing_and_other_family_is_not_overwritten(self):
        self.assertEqual(product(props={'Материал корпуса':'—'}).values['material'],'brass')
        p=product(props={'Тип датчика':'Ёмкостный'})
        self.assertEqual(p.family,'capacitive');self.assertNotIn('diameter',p.values)

    def test_original_document_is_immutable_and_exports_keep_evidence(self):
        record={'article':'PS2-18M68-8N11-K','source':'megak','_specifications':{'attributes':[]}}
        original=deepcopy(record);p=normalize_sensor(record)
        self.assertEqual(record,original)
        rows=matching_export(p,evaluate(p,sensor(sn=8.,mount='non-flush')))
        row=next(r for r in rows if r['Характеристика']=='Длина корпуса, мм')
        self.assertIn('Обозначение МЕГА-К',row['Источник конкурента'])
        self.assertEqual(row['Источник расшифровки'],SOURCE_URL)
        wb=load_workbook(BytesIO(xlsx_bytes(rows)))
        self.assertIn('Источник конкурента',[c.value for c in wb.active[1]])

    def test_full_product_tab_and_summary_are_merged_and_deduplicated(self):
        html='''<ul id="properties"><li><div class="properties-item-row"><div class="properties-item-name">Материал корпуса</div><div class="properties-item-value">латунь</div></div></li></ul>
        <div><span class="details-param-name">Материал корпуса:</span><span class="details-param-value">латунь</span></div>
        <aside><div class="properties-item-row"><div class="properties-item-name">Чужой товар</div><div class="properties-item-value">не брать</div></div></aside>'''
        self.assertEqual(attributes('megak',BeautifulSoup(html,'html.parser')),[{'name':'Материал корпуса','value':'латунь','group':''}])

    def test_megak_length_label_does_not_confuse_cable_or_packaging(self):
        p=product(props={'Длина':'70 мм','Длина кабеля':'2 м','Размеры':'120x90x55 мм'})
        self.assertEqual(p.values['length'],70.)
        self.assertEqual(p.values['diameter'],18.)
        self.assertEqual(p.decoding['fields']['length']['application'],'conflict')

    def test_archived_real_product_tab_preserves_card_load_and_frequency(self):
        path=Path(__file__).resolve().parents[1]/'docs/prices/audit/raw/megak_extra_0.txt'
        if not path.exists():self.skipTest('Optional archived source is not part of the release')
        attrs=attributes('megak',BeautifulSoup(path.read_text(encoding='utf-8'),'html.parser'))
        p=normalize_sensor({'manufacturer':'МЕГА-К','article':'PS2-36M70-12B11-K','_specifications':{'attributes':attrs}})
        self.assertEqual((p.values['diameter'],p.values['length'],p.values['sn']),(36.,70.,12.))
        self.assertEqual((p.values['frequency'],p.values['load']),(300.,300.))
        self.assertEqual(p.decoding['fields']['load']['application'],'card_priority')
        self.assertGreater(len(attrs),30)


if __name__=='__main__':unittest.main()

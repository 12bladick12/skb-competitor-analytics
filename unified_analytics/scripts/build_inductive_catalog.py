"""Build the app's reference catalog from the supplied TDSheet workbook.

Run from the application root: python scripts/build_inductive_catalog.py PATH.xlsx
Only active inductive products and matching characteristics are included.
"""
import argparse
import hashlib
import gzip
import json
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
from price_monitor.matching_normalize import ALIASES, PREFIXES, key


def build(source:Path):
    allowed={key(x) for values in ALIASES.values() for x in values}
    prefixes=[key(x) for x in PREFIXES.values()]+[key('_Дополнительные характеристики'),key('_Специальное назначение')]
    wb=load_workbook(source,read_only=True,data_only=True)
    sheet=wb['TDSheet'];products=[];category='';current=None
    for row_no,row in enumerate(sheet.iter_rows(values_only=True),1):
        row=tuple(row)+(None,)*8
        if row_no<=4:continue
        if row[6] is not None and str(row[7]).strip() in ('Да','Нет'):
            current=None
            if row[7]=='Нет' and 'индуктив' in category.casefold():
                current={'model':str(row[0]),'article':str(row[5] or ''),'code':str(row[6]),'category':category,'row':row_no,'props':{}}
                products.append(current)
        elif row[0] is not None and row[5] is None and row[6] is None:
            category=str(row[0]);current=None
        elif current is not None and row[0] is not None and row[5] is not None:
            name=key(row[0])
            if name in allowed or any(name.startswith(prefix) for prefix in prefixes):current['props'][str(row[0])]=row[5]
    wb.close()
    if not products:raise ValueError('Не найдены активные индуктивные датчики на листе TDSheet')
    payload={'schema_version':1,'snapshot_date':'2026-09-29','source_modified_at':'2026-08-28',
             'source_url':'https://docs.google.com/spreadsheets/d/1vmRnTUjkkR06a3mz69ze0wbrv391JQ5SyHY6HgZzk3Y/edit',
             'source_sheet':'TDSheet','source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'products':products}
    dest=ROOT/'price_monitor/assets/our_inductive_catalog.json.gz';dest.parent.mkdir(exist_ok=True)
    dest.write_bytes(gzip.compress(json.dumps(payload,ensure_ascii=False,separators=(',',':')).encode('utf-8'),mtime=0))
    print(f'Active inductive products: {len(products)}; snapshot bytes: {dest.stat().st_size}')
    return dest

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source',type=Path)
    build(parser.parse_args().source)

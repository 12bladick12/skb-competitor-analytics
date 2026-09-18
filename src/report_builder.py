from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell

from .models import AppSettings, CompetitorConfig, ReportPeriod
from .storage import Storage
from .utils import normalize_text, report_period_key


SIGNIFICANCE = {
    "новая продукция": 1,
    "изменение позиционирования": 2,
    "акция": 3,
    "технический материал": 4,
    "коммерческое предложение": 5,
    "публикация в Telegram": 6,
    "новость": 7,
}


class ReportBuilder:
    """Create a compact, evidence-first monthly management note."""

    def __init__(self, storage: Storage, settings: AppSettings):
        self.storage = storage
        self.settings = settings
        self.primary = RGBColor.from_string(settings.report_primary_color)
        self.muted = RGBColor.from_string(settings.report_muted_color)

    def build(self, competitors: list[CompetitorConfig], period: ReportPeriod) -> Path:
        from .verified_report import build_verified_report
        return build_verified_report(self, competitors, period)

    def _setup_document(self, doc: Document) -> None:
        section = doc.sections[0]
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        styles = doc.styles
        styles["Normal"].font.name = self.settings.report_font_name
        styles["Normal"].font.size = Pt(9.5)
        styles["Normal"].paragraph_format.space_after = Pt(5)
        for style_name, size in (("Heading 1", 14), ("Heading 2", 11)):
            style = styles[style_name]
            style.font.name = self.settings.report_font_name
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = self.primary

    def _save_document(self, doc: Document, path: Path) -> Path:
        temporary = path.with_name(f".{path.stem}.tmp{path.suffix}")
        try:
            doc.save(temporary)
            temporary.replace(path)
            return path
        except PermissionError as exc:
            if temporary.exists():
                temporary.unlink()
            raise RuntimeError(f"Закройте открытый файл отчёта и повторите запуск: {path.name}") from exc
        except Exception:
            if temporary.exists():
                temporary.unlink()
            raise

    def _title(self, doc: Document, competitors: list[CompetitorConfig], period: ReportPeriod) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Аналитическая записка по конкурентам")
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = self.primary
        for text in (
            f"Период анализа: {period.label}",
            f"Дата формирования: {datetime.now():%d.%m.%Y}",
            "Конкуренты: " + ", ".join(item.name for item in competitors),
        ):
            paragraph = doc.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.runs[0].font.color.rgb = self.muted

    def _signature_block(self, doc: Document) -> None:
        for role, name in (
            ("Подготовил", "Селезнев В.Е."),
            ("Согласовал(а)", "Никанкина А.С."),
        ):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.space_after = Pt(1)
            role_run = paragraph.add_run(role)
            role_run.bold = True
            paragraph.add_run("  ____________________  ")
            paragraph.add_run(name)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def _summary(self, doc: Document, events: list[dict], quality: dict[str, int]) -> None:
        self._heading(doc, "1. Ключевые выводы")
        if not events:
            doc.add_paragraph("За период не выявлено подтверждённых новостных материалов, пригодных для управленческого анализа.")
        else:
            active = Counter(item["competitor_name"] for item in events)
            categories = Counter(item["category"] for item in events)
            bullets = [
                f"Подтверждено материалов за период: {len(events)}.",
                f"Наиболее активный конкурент: {active.most_common(1)[0][0]} ({active.most_common(1)[0][1]} материалов).",
                "Основные типы сигналов: " + ", ".join(f"{name} — {count}" for name, count in categories.most_common(3)) + ".",
            ]
            for bullet in bullets:
                doc.add_paragraph(bullet, style="List Bullet")
        doc.add_paragraph(
            f"Контроль качества: подтверждено {quality['verified']}; без подтверждённой даты {quality['unconfirmed']}; "
            f"исключено как не-статья {quality['excluded']}.",
        ).runs[0].font.color.rgb = self.muted

    def _signals(self, doc: Document, events: list[dict], total: int) -> None:
        self._heading(doc, "3. Подтверждённые сигналы")
        if not events:
            doc.add_paragraph("Значимые подтверждённые сигналы отсутствуют.")
            return
        if total > len(events):
            doc.add_paragraph(
                f"Показаны {len(events)} наиболее значимых материалов из {total}; "
                "остальные учтены в итоговых показателях."
            )
        rows = [["Дата", "Конкурент", "Сигнал", "Практическое значение", "Источник"]]
        for item in events:
            rows.append([
                self._date(item["published_at"]),
                item["competitor_name"],
                self._signal_text(item),
                self._implication(item),
                item["url"],
            ])
        self._table(doc, rows)

    def _risks_and_opportunities(self, doc: Document, events: list[dict]) -> None:
        self._heading(doc, "4. Риски и возможности")
        significant = [item for item in events if item["category"] in SIGNIFICANCE and item["category"] != "новость"]
        if not significant:
            doc.add_paragraph("Подтверждённые риски и возможности, требующие отдельной реакции, не выявлены.")
            return
        rows = [["Наблюдение", "Что это означает", "Рекомендуемая реакция"]]
        for item in significant[:5]:
            rows.append([
                f"{item['competitor_name']}: {item['category']}",
                self._implication(item),
                self._action_for(item),
            ])
        self._table(doc, rows)

    def _site_activity(self, doc: Document, events: list[dict], watches: list[dict]) -> None:
        self._heading(doc, "2. Обновления сайтов")
        if not events:
            doc.add_paragraph("За период новые изменения отслеживаемых фрагментов сайтов не зафиксированы.")
        else:
            totals: dict[str, Counter] = {}
            for item in events:
                counter = totals.setdefault(str(item["competitor_name"]), Counter())
                counter["all"] += 1
                if item.get("event_kind") == "новость":
                    counter["confirmed"] += int(item.get("status") == "confirmed")
                else:
                    counter["page_changes"] += 1
            summary_rows = [["Конкурент", "Подтв. новости", "Изм. страниц", "Всего"]]
            for competitor, counter in sorted(totals.items()):
                summary_rows.append([competitor, str(counter["confirmed"]), str(counter["page_changes"]), str(counter["all"])])
            self._table(doc, summary_rows)

            visible = [item for item in events if item.get("event_kind") != "новость" or item.get("status") == "confirmed"]
            if not visible:
                doc.add_paragraph("Существенных обновлений с подтверждённой датой новости не выявлено; непроверенные новости не включены в документ.")
            else:
                rows = [["Дата", "Конкурент", "Тип", "Контекст и diff", "Источник"]]
                for item in visible[:20]:
                    evidence = self._short(str(item.get("diff_text") or ""), 210)
                    context = self._short(str(item.get("summary") or item.get("title") or "Обновление страницы"), 155)
                    if evidence:
                        context = f"{context}\n{evidence}"
                    source = str(item.get("article_url") or item.get("source_url") or "")
                    rows.append([
                        self._date(item.get("published_at") or item.get("detected_at")),
                        str(item["competitor_name"]),
                        str(item.get("event_kind") or "изменение страницы"),
                        context,
                        source,
                    ])
                self._table(doc, rows)
                if len(visible) > 20:
                    doc.add_paragraph(f"Ещё {len(visible) - 20} обновлений сохранены в журнале site_change_events.").runs[0].font.color.rgb = self.muted

        unavailable = [watch for watch in watches if watch.get("status") == "error"]
        if unavailable:
            doc.add_paragraph("Недоступные или требующие проверки наблюдения:")
            rows = [["Конкурент", "Раздел", "Причина", "Источник"]]
            for watch in unavailable[:10]:
                rows.append([
                    str(watch["competitor_name"]),
                    str(watch.get("title") or watch.get("source_type") or "Страница"),
                    self._short(str(watch.get("last_error") or "Ошибка проверки"), 140),
                    str(watch.get("url") or ""),
                ])
            self._table(doc, rows)

    def _telegram_appendix(self, doc: Document, events: list[dict]) -> None:
        posts = [item for item in events if item.get("source_type") == "telegram"]
        if not posts:
            return
        self._heading(doc, "Приложение. Публикации Teko в Telegram")
        doc.add_paragraph(
            f"За период собрано {len(posts)} публикаций. Дата подтверждена временной меткой Telegram; "
            "в основной аналитике учитываются только наиболее значимые сигналы."
        )
        rows = [["Дата", "Краткое содержание", "Источник"]]
        for item in sorted(posts, key=lambda value: str(value["published_at"]), reverse=True)[:20]:
            rows.append([
                self._date(item["published_at"]),
                self._short(item.get("summary") or item.get("title") or "Публикация Telegram", 170),
                item["url"],
            ])
        self._table(doc, rows)
        if len(posts) > 20:
            doc.add_paragraph(f"Ещё {len(posts) - 20} публикаций сохранены в базе данных.").runs[0].font.color.rgb = self.muted

    @staticmethod
    def _visible_events(events: list[dict], limit: int) -> list[dict]:
        site_events = [item for item in events if item.get("source_type") != "telegram"]
        telegram = [item for item in events if item.get("source_type") == "telegram"][:3]
        site_limit = max(0, limit - len(telegram))

        selected_site: list[dict] = []
        selected_ids: set[int] = set()
        seen_competitors: set[str] = set()
        for item in site_events:
            competitor = str(item.get("competitor_name") or "")
            if competitor in seen_competitors:
                continue
            selected_site.append(item)
            selected_ids.add(id(item))
            seen_competitors.add(competitor)
            if len(selected_site) >= site_limit:
                break
        for item in site_events:
            if len(selected_site) >= site_limit:
                break
            if id(item) not in selected_ids:
                selected_site.append(item)
                selected_ids.add(id(item))

        selected = [*selected_site, *telegram]
        return sorted(selected, key=lambda item: str(item["published_at"]), reverse=True)

    def _verified_events(self, period: ReportPeriod) -> list[dict]:
        rows = self.storage.fetchall(
            """
            SELECT p.*, c.name AS competitor_name
            FROM pages p JOIN competitors c ON c.id = p.competitor_id
            WHERE p.published_at BETWEEN ? AND ?
              AND p.content_kind = 'article'
              AND (
                    (p.source_type = 'news' AND p.publication_date_status = 'verified')
                    OR (p.source_type IN ('promotions', 'docs') AND p.published_at IS NOT NULL)
                  )
              AND COALESCE(TRIM(p.text), '') != ''
            ORDER BY p.published_at DESC, p.id DESC
            """,
            (period.start.isoformat(timespec="seconds"), period.end.isoformat(timespec="seconds")),
        )
        events = []
        seen: set[tuple[int, str]] = set()
        for row in rows:
            item = dict(row)
            key = (int(item["competitor_id"]), str(item["url"]).split("#", 1)[0].split("?", 1)[0].rstrip("/").lower())
            if key in seen:
                continue
            seen.add(key)
            item["category"] = item["classification_category"] or "новость"
            item["summary"] = item["short_summary"] or item["text"] or ""
            events.append(item)
        events.sort(key=lambda item: item["published_at"], reverse=True)
        return sorted(events, key=lambda item: SIGNIFICANCE.get(item["category"], 99))

    def _telegram_events(self, period: ReportPeriod, competitors: list[CompetitorConfig]) -> list[dict]:
        channel_names: dict[str, str] = {}
        for competitor in competitors:
            if not competitor.telegram or not competitor.telegram.enabled:
                continue
            channel = competitor.telegram.url.rstrip("/").lower()
            channel_names[channel] = competitor.name
            channel_names[channel.rsplit("/", 1)[-1].lstrip("@")] = competitor.name

        rows = self.storage.fetchall(
            """
            SELECT * FROM telegram_posts
            WHERE published_at BETWEEN ? AND ?
              AND COALESCE(TRIM(text), '') != ''
            ORDER BY published_at DESC, id DESC
            """,
            (period.start.isoformat(timespec="seconds"), period.end.isoformat(timespec="seconds")),
        )
        events: list[dict] = []
        seen: set[str] = set()
        default_name = next(iter(channel_names.values()), "Telegram")
        for row in rows:
            item = dict(row)
            key = str(item.get("url") or item.get("post_id") or item["id"])
            if key in seen:
                continue
            seen.add(key)
            channel = str(item.get("channel") or "").rstrip("/").lower()
            competitor_name = channel_names.get(channel) or channel_names.get(channel.rsplit("/", 1)[-1].lstrip("@")) or default_name
            post_text = normalize_text(str(item.get("text") or ""))
            title = post_text.split(". ", 1)[0]
            events.append({
                **item,
                "competitor_name": competitor_name,
                "title": self._short(title, 105),
                "summary": post_text,
                "category": "публикация в Telegram",
                "source_type": "telegram",
            })
        return events

    def _unconfirmed_news(self, period: ReportPeriod) -> list[dict]:
        rows = self._canonical_news_rows(period)
        return [
            item
            for item in rows
            if item["publication_date_status"] != "verified" or item["content_kind"] != "article"
        ]

    def _canonical_news_rows(self, period: ReportPeriod) -> list[dict]:
        rows = self.storage.fetchall(
            """
            SELECT p.*, c.name AS competitor_name, a.reason
            FROM pages p JOIN competitors c ON c.id = p.competitor_id
            LEFT JOIN page_date_audits a ON a.id = (
                SELECT latest.id FROM page_date_audits latest WHERE latest.page_id = p.id ORDER BY latest.id DESC LIMIT 1
            )
            WHERE p.source_type = 'news'
              AND p.discovered_at BETWEEN ? AND ?
            ORDER BY p.discovered_at DESC, p.id DESC
            """,
            (period.start.isoformat(timespec="seconds"), period.end.isoformat(timespec="seconds")),
        )
        result_by_key: dict[tuple[int, str], dict] = {}
        for row in rows:
            item = dict(row)
            key = (int(item["competitor_id"]), str(item["url"]).split("#", 1)[0].split("?", 1)[0].rstrip("/").lower())
            previous = result_by_key.get(key)
            if previous is None or self._canonical_news_score(item) > self._canonical_news_score(previous):
                result_by_key[key] = item
        return sorted(result_by_key.values(), key=lambda item: str(item.get("discovered_at") or ""), reverse=True)

    @staticmethod
    def _canonical_news_score(item: dict) -> tuple[int, int, str, int]:
        return (
            int(item.get("publication_date_status") == "verified" and item.get("content_kind") == "article"),
            int(item.get("content_kind") == "article"),
            str(item.get("checked_at") or item.get("discovered_at") or ""),
            int(item.get("id") or 0),
        )

    def _quality_summary(self, period: ReportPeriod, events: list[dict]) -> dict[str, int]:
        rows = self._canonical_news_rows(period)
        excluded = sum(item["content_kind"] != "article" for item in rows)
        missing_articles = sum(
            item["content_kind"] == "article" and item["publication_date_status"] != "verified"
            for item in rows
        )
        return {"verified": len(events), "unconfirmed": missing_articles, "excluded": excluded}

    def _heading(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph(text, style="Heading 1")
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)

    def _table(self, doc: Document, rows: list[list[str]]) -> None:
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = ([0.9,1.0,4.4,0.8] if len(rows[0])==4 and 'Материал' in rows[0][2]
                  else [1.05,1.8,2.5,1.75] if len(rows[0])==4 else [7.1/len(rows[0])]*len(rows[0]))
        for column,width in zip(table.columns,widths):
            column.width=Inches(width)
        for row_index, values in enumerate(rows):
            cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
            properties = table.rows[row_index]._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if row_index == 0:
                properties.append(OxmlElement('w:tblHeader'))
            for cell, value,width in zip(cells, values,widths):
                cell.width=Inches(width)
                self._set_cell_value(cell, value)
                self._format_cell(cell, header=row_index == 0)

    def _set_cell_value(self, cell: _Cell, value: str) -> None:
        paragraph = cell.paragraphs[0]
        if isinstance(value,tuple):
            self._add_hyperlink(paragraph,value[1],value[0])
        elif value.startswith("http://") or value.startswith("https://"):
            self._add_hyperlink(paragraph, value, "Открыть")
        else:
            paragraph.add_run(value)

    def _add_hyperlink(self, paragraph, url: str, text: str) -> None:
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(color)
        properties.append(underline)
        run.append(properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        run.append(text_element)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def _format_cell(self, cell: _Cell, header: bool = False) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if header:
            properties = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), self.settings.report_primary_color)
            properties.append(shading)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(1)
            for run in paragraph.runs:
                run.font.name = self.settings.report_font_name
                run.font.size = Pt(9)
                if header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

    def _signal_text(self, item: dict) -> str:
        return f"{item['category'].capitalize()}: «{self._short(item['title'], 105)}»"

    def _implication(self, item: dict) -> str:
        category = item["category"]
        if category == "новая продукция":
            return "Сравнить характеристики с линейкой СКБ «Индукция»."
        if category == "акция":
            return "Оценить коммерческое предложение и риск ценового давления."
        if category == "изменение позиционирования":
            return "Сверить собственные аргументы продаж и коммуникационные акценты."
        if category == "технический материал":
            return "Проверить полноту собственных паспортов, инструкций и схем."
        if category == "коммерческое предложение":
            return "Оценить применимость аргумента в сравнительных материалах."
        if category == "публикация в Telegram":
            return "Учесть оперативный коммуникационный сигнал ТЕКО и проверить его влияние на клиентов и продажи."
        title = normalize_text(str(item.get("title") or "")).lower()
        if any(marker in title for marker in ("upgrade", "expand", "new-generation", "smartlight", "sensor series")):
            return "Сопоставить заявленное обновление с собственной продуктовой линейкой и аргументами продаж."
        if any(marker in title for marker in ("webinar", "how to choose", "guide", "measurement")):
            return "Сравнить охват темы с собственными техническими материалами и консультационными сценариями."
        if any(marker in title for marker in ("invitation", "exhibition", "taipei", "expo")):
            return "Отслеживать анонсы продуктов и отраслевые акценты конкурента вокруг мероприятия."
        if any(marker in title for marker in ("designated", "award", "member", "club")):
            return "Учесть усиление репутационного позиционирования конкурента в коммуникациях с рынком."
        if any(marker in title for marker in ("temperature", "range", "диапазон")):
            return "Проверить сопоставимость рабочих диапазонов и актуальность сравнительных материалов."
        return "Подтверждён факт публикационной активности; дополнительный бизнес-эффект по материалу не установлен."

    def _action_for(self, item: dict) -> str:
        if item["category"] == "новая продукция":
            return "Назначить сравнение характеристик и целевого сегмента."
        if item["category"] == "акция":
            return "Проверить цены, условия и альтернативные предложения."
        if item["category"] == "технический материал":
            return "Обновить недостающие материалы по сопоставимым изделиям."
        return "Уточнить аргументы продаж и коммуникации."

    @staticmethod
    def _date(value: str | None) -> str:
        if not value:
            return "—"
        try:
            return datetime.fromisoformat(value).strftime("%d.%m.%Y")
        except ValueError:
            return value[:10]

    @staticmethod
    def _short(value: str, limit: int) -> str:
        clean = normalize_text(value)
        return clean if len(clean) <= limit else clean[: limit - 1].rstrip() + "…"

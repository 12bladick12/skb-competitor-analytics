from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document

from .briefs import russian_brief
from .evidence_store import EvidenceStore
from .utils import report_period_key
from .verified_monitor import sources_for

LABELS = {'success':'Проверен', 'partial':'Сбор неполный', 'error':'Ошибка проверки',
          'not_configured':'Канал не включён в мониторинг'}


def build_verified_report(builder, competitors, period):
    from .verified_monitor import VerifiedMonitor
    verifier=VerifiedMonitor(builder.storage,builder.settings,competitors)
    try:
        verifier.revalidate_saved_events(period)
    finally:
        verifier.close()
    from .report_content import prepare_report_data
    events, checks, link_checks = prepare_report_data(builder.storage, competitors, period)
    lookup = {(x['competitor_code'],x['kind'],x['url']):x for x in checks}
    rows = [['Конкурент','Новости сайтов','Продукция и предложения','Telegram']]
    coverage_links = []
    incomplete = False
    incomplete |= any(link_checks.get(e['id'],{}).get('status')=='unavailable' for e in events)
    for competitor in competitors:
        groups = []
        row = [competitor.name]
        for kinds in [('news',),('products','promotions'),('telegram',)]:
            sources = [(s['kind'],s['url']) for s in sources_for(competitor) if s['kind'] in kinds]
            if kinds == ('telegram',):
                sources = [('telegram',competitor.telegram.url if competitor.telegram and competitor.telegram.enabled else '')]
            descriptions, links = [], []
            for kind,url in sources:
                check = lookup.get((competitor.code,kind,url))
                status = check['status'] if check else ('not_configured' if not url else 'partial')
                incomplete |= status in {'partial','error'}
                description = LABELS[status]
                source_definition = next((s for s in sources_for(competitor) if s['url']==url),{})
                if status=='success' and source_definition.get('mode')=='catalogue':
                    description='Текущий каталог проверен. История до первого снимка не подтверждена'
                reason = (check or {}).get('reason','')
                if status in {'partial','error'}:
                    if 'limit' in reason:
                        description += ': обход не завершён'
                    elif 'date' in reason:
                        description += ': часть дат не подтверждена'
                    elif 'cards' in reason or 'content' in reason:
                        description += ': содержимое получено не полностью'
                success = (check or {}).get('last_success_at')
                if success:
                    description += '\nУспешно: '+datetime.fromisoformat(success).strftime('%d.%m %H:%M')+' UTC'
                if kind == 'promotions':
                    description = 'Предложения: '+description
                descriptions.append(description)
                if url:
                    links.append(({'news':'Новости','products':'Продукция','promotions':'Предложения','telegram':'Канал'}[kind],url))
            row.append('\n'.join(descriptions) if descriptions else 'Продуктовые анонсы проверяются в новостях')
            groups.append(links)
        rows.append(row)
        coverage_links.append(groups)
    doc = Document()
    builder._setup_document(doc)
    builder._signature_block(doc)
    builder._title(doc,competitors,period)
    if incomplete:
        doc.add_paragraph('СБОР НЕПОЛНЫЙ. Ниже приведены только подтверждённые материалы. Отсутствие записей по недоступному источнику не означает отсутствие новостей.').runs[0].bold=True
    counts = Counter(e['kind'] for e in events)
    doc.add_paragraph(f"Сводка: новости сайтов — {counts['news']}; публикации Telegram — {counts['telegram']}; обновления продукции и предложений — {counts['products']}. Всего записей: {len(events)}.")
    doc.add_paragraph('Один материал учитывается один раз внутри своего раздела. Повторная публикация в Telegram показана отдельно как публикация канала. Заявления производителей передаются по первоисточнику.')
    doc.add_paragraph('Первое наблюдение товарного раздела создаёт базу сравнения и не подтверждает появление продукции в прошлом. Отсутствие исторических изменений в отчёте не означает, что их не было.')
    builder._heading(doc,'Полнота проверки источников')
    builder._table(doc,rows)
    for row,groups in zip(doc.tables[-1].rows[1:],coverage_links):
        for cell,links in zip(row.cells[1:],groups):
            for label,url in links:
                paragraph=cell.add_paragraph()
                from docx.shared import Pt
                paragraph.paragraph_format.space_after=Pt(0)
                builder._add_hyperlink(paragraph,url,label)
    for kind,heading in [('news','1. Новости сайтов'),('telegram','2. Новости Telegram-каналов'),('products','3. Обновления продукции и предложений')]:
        builder._heading(doc,heading)
        group = [e for e in events if e['kind']==kind]
        if not group:
            doc.add_paragraph('Подтверждённых материалов за период в собранных данных нет. Полнота проверки указана выше.')
            continue
        table = [['Дата','Конкурент','Материал и краткое содержание','Источник']]
        for event in group:
            evidence = json.loads(event['evidence_json'])
            date = builder._date(event['published_at'] or event['detected_at'])
            if event['date_basis']=='observed_change':
                from datetime import timezone
                from .telegram_history import LOCAL
                date = datetime.fromisoformat(event['detected_at']).replace(tzinfo=timezone.utc).astimezone(LOCAL).strftime('%d.%m.%Y')+' (обнаружено, Екатеринбург)'
            elif event['date_basis']=='announcement_month':
                date = datetime.fromisoformat(event['published_at']).strftime('%m.%Y')+' (месяц анонса)'
            content = russian_brief(event['original_text'],builder.settings.translation_enabled)
            title = event['title']
            if event['kind']=='telegram' and len(title)>95:
                title = title[:95].rsplit(' ',1)[0].rstrip('.,:;')+'…'
            link = evidence.get('article_url') or event['url']
            if link_checks.get(event['id'],{}).get('status')=='unavailable':
                import shutil
                copy_dir=builder.settings.reports_dir/'evidence'
                copy_dir.mkdir(parents=True,exist_ok=True)
                copy=copy_dir/f"event_{event['id']}.html"
                shutil.copyfile(evidence['snapshots'][0],copy)
                link=('Сохранённая копия',copy.resolve().as_uri())
                content += '\nОригинальная страница сейчас недоступна; материал подтверждён сохранённым снимком.'
            table.append([date,event['competitor_name'],title+'\n'+content,link])
        builder._table(doc,table)
    path = builder.settings.reports_dir / f"Аналитическая_записка_конкуренты_{report_period_key(period).replace('-','_')}.docx"
    doc.core_properties.title='Аналитическая записка по конкурентам'
    doc.core_properties.subject=f'Проверенные материалы: {period.label}'
    doc.core_properties.author='СКБ «Индукция»'
    builder.settings.reports_dir.mkdir(parents=True,exist_ok=True)
    output = builder._save_document(doc,path)
    audit = {'period':period.label,'incomplete':incomplete,'counts':dict(counts),'sources':checks,
             'events':[{'id':e['id'],'competitor':e['competitor_code'],'kind':e['kind'],'url':e['url'],'link_check':link_checks.get(e['id']),
                        'date':e['published_at'],'date_basis':e['date_basis'],'evidence':json.loads(e['evidence_json'])} for e in events]}
    audit_path = builder.settings.processed_dir / f'verification_{report_period_key(period)}.json'
    audit_path.parent.mkdir(parents=True,exist_ok=True)
    audit_path.write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    return output

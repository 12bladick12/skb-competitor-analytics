"""Decision regressions agreed with the user, plus parsing/persistence boundaries."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import json
import tempfile
import unittest
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from price_monitor.matching import evaluate, rank, Matcher, load_catalog, matching_export
from price_monitor.matching_normalize import Sensor, normalize_sensor, interval
from price_monitor.exchange import xlsx_bytes
from price_monitor.storage import Store
from price_monitor.library import Library


def sensor(name='Исходный',**changes):
    values=dict(body_type='threaded',diameter=18.,pitch=1.,output='PNP',function='NO',mount='flush',voltage_type='DC',vmin=10.,vmax=30.,
                sn=5.,tmin=-45.,tmax=85.,material='brass',connection='cable',ip=('67',),load=200.,frequency=500.,length=50.)
    values.update(changes)
    return Sensor(name,name,'Индуктивные датчики',family='inductive',values=values)


class AgreedDecisions(unittest.TestCase):
    def test_cold_example_v_before_a_before_b(self):
        reference=sensor()
        candidates=[sensor('А',length=65.),sensor('Б',tmin=-25.),sensor('В',tmin=-60.,tmax=70.,length=55.)]
        matches=rank([evaluate(reference,c,'cold') for c in candidates])
        self.assertEqual([m.candidate.model for m in matches],['В','А','Б'])
        self.assertEqual([m.status for m in matches],['direct','direct','close'])

    def test_cold_rule_does_not_apply_to_general_profile(self):
        result=evaluate(sensor(),sensor('В',tmin=-60.,tmax=70.,length=55.),'general')
        self.assertEqual(result.status,'close')
        self.assertEqual(next(f for f in result.fields if f['key']=='tmax')['outcome'],'difference')

    def test_exact_temperature_before_surplus_only_when_otherwise_equal(self):
        results=rank([evaluate(sensor(),sensor('Запас',tmin=-60.),'cold'),evaluate(sensor(),sensor('Точное'),'cold')])
        self.assertEqual(results[0].candidate.model,'Точное')

    def test_sn_symmetry_and_no_unagreed_tolerance(self):
        results=rank([evaluate(sensor(),sensor('Меньше',sn=4.)),evaluate(sensor(),sensor('Больше',sn=6.))])
        self.assertEqual({m.priority for m in results},{1})
        self.assertEqual({m.status for m in results},{'close'})
        self.assertEqual(evaluate(sensor(),sensor('Очень близко',sn=5.01)).status,'close')

    def test_material_and_connection_equal_priority(self):
        result=rank([evaluate(sensor(),sensor('Материал',material='stainless')),evaluate(sensor(),sensor('Подключение',connection='connector'))])
        self.assertEqual({m.priority for m in result},{1})
        self.assertEqual({m.status for m in result},{'close'})

    def test_incomparable_important_losses_keep_both_options(self):
        result=rank([evaluate(sensor(),sensor('Sn',sn=4.,length=51.)),evaluate(sensor(),sensor('Материал',material='plastic',length=80.))])
        self.assertEqual({m.priority for m in result},{1})

    def test_diameter_and_pitch_cannot_be_compensated(self):
        for changes in ({'diameter':12.},{'pitch':1.5},{'output':'NPN'},{'mount':'non-flush'},{'function':'NC'}):
            self.assertEqual(evaluate(sensor(),sensor('Лучше',tmin=-80.,frequency=2000.,**changes),'cold').status,'incompatible')

    def test_unknown_pitch_is_not_a_direct_match(self):
        result=evaluate(sensor(pitch=None),sensor())
        self.assertEqual(result.status,'review')
        self.assertIn('pitch',[f['key'] for f in result.missing])

    def test_length_secondary_unless_mounting_limit_given(self):
        self.assertEqual(evaluate(sensor(),sensor(length=80.)).status,'direct')
        self.assertEqual(evaluate(sensor(),sensor(length=80.),max_length=70).status,'incompatible')

    def test_load_improvement_and_reversal_are_different(self):
        self.assertEqual(evaluate(sensor(),sensor(load=300.)).status,'direct')
        self.assertEqual(evaluate(sensor(load=300.),sensor()).status,'close')

    def test_ip_is_not_numeric_order(self):
        self.assertEqual(evaluate(sensor(ip=('67',)),sensor(ip=('65',))).status,'close')
        self.assertEqual(evaluate(sensor(ip=('65',)),sensor(ip=('67',))).status,'close')
        self.assertEqual(evaluate(sensor(ip=('65',)),sensor(ip=('65','67'))).status,'direct')

    def test_supply_range_and_current_kind(self):
        self.assertEqual(evaluate(sensor(),sensor(vmin=5.,vmax=36.)).status,'direct')
        self.assertEqual(evaluate(sensor(),sensor(vmin=12.)).status,'incompatible')
        self.assertEqual(evaluate(sensor(voltage_type='AC/DC'),sensor(voltage_type='DC')).status,'incompatible')

    def test_special_flags_can_coexist_but_do_not_get_standard_approval(self):
        reference=replace(sensor(),special=('namur','pressure'))
        self.assertEqual(evaluate(reference,replace(sensor(),special=('namur','pressure'))).status,'review')
        self.assertEqual(evaluate(reference,replace(sensor(),special=('namur',))).status,'incompatible')

    def test_missing_family_and_other_family(self):
        self.assertEqual(evaluate(replace(sensor(),family=None),sensor()).status,'review')
        self.assertEqual(evaluate(replace(sensor(),family='capacitive'),sensor()).status,'unsupported')


class NormalizationChecks(unittest.TestCase):
    def test_range_signs_and_units(self):
        self.assertEqual(interval('10-30[DC]'),(10.,30.))
        self.assertEqual(interval('-45…+85 °С'),(-45.,85.))
        self.assertIsNone(interval('24 [DC]'))
        p=normalize_sensor({'category':'Индуктивные','props':{'Макс. ток нагрузки, А:':'0,2','Частота переключения, Гц':'1 кГц'}})
        self.assertEqual(p.values['load'],200.)
        self.assertEqual(p.values['frequency'],1000.)

    def test_m18_without_pitch_and_sa_do_not_fill_critical_values(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Размер резьбы корпуса:':'M18','Рабочее расстояние переключения, Sa, мм':'4,1'}})
        self.assertEqual(p.values['diameter'],18.)
        self.assertNotIn('pitch',p.values);self.assertNotIn('sn',p.values)

    def test_conflicting_structured_fields_remain_unknown(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Типоразмер корпуса, мм':'М18x1','Диаметр резьбового корпуса':'M12x1'}})
        self.assertNotIn('diameter',p.values);self.assertIn('diameter',p.conflicts)

    def test_family_not_inferred_from_electrical_output(self):
        p=normalize_sensor({'title':'Датчик положения','props':{'Схема выхода':'NPN'}})
        self.assertIsNone(p.family);self.assertNotIn('namur',p.special)

    def test_temperature_value_does_not_select_cold_profile(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Рабочая температура окружающей среды, °С':'-45…+85'}})
        self.assertEqual(p.profile,'general')

    def test_explicit_cold_and_combined_special_categories(self):
        p=normalize_sensor({'category':'Индуктивные датчики высокого давления стандарта NAMUR','props':{'_Дополнительные характеристики':'для низких температур'}})
        self.assertEqual(p.special,('namur','pressure'));self.assertEqual(p.profile,'cold')

    def test_document_voltage_current_and_speed_fields_not_conflated(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Аналоговый выход по напряжению, В':'0…10','Диапазон измерения частоты, Гц':'2…50','Номинальное напряжение питания, В':'24'}})
        self.assertNotIn('load',p.values);self.assertNotIn('frequency',p.values);self.assertNotIn('vmin',p.values)
        self.assertEqual(p.special,('analog','speed'))

    def test_special_properties_do_not_receive_standard_approval(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Максимальное рабочее давление, бар':'500','Маркировка взрывозащиты':'1Ex ia IIC T6','Аналоговый выход по току, мА':'—'}})
        self.assertEqual(p.special,('ex','pressure'))

    def test_explicit_other_family_overrides_broad_category(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Тип датчика':'Ёмкостный'}})
        self.assertEqual(p.family,'capacitive')

    def test_packaging_not_housing_and_connector_not_body(self):
        p=normalize_sensor({'category':'Индуктивные','props':{'Размеры упаковки':'M30x1,5','Соединение':'Разъем M12'}})
        self.assertNotIn('diameter',p.values);self.assertEqual(p.values['connection'],'connector')


class IntegrationChecks(unittest.TestCase):
    def test_catalog_uses_active_provided_products(self):
        payload,matcher=load_catalog()
        self.assertEqual(len(matcher.products),1324)
        self.assertEqual(len({p.id for p in matcher.products}),1324)
        self.assertTrue(all(p.family=='inductive' for p in matcher.products))
        p=matcher.resolve('ИВ05-NO-PNP(Л63)')
        self.assertEqual(p.values['sn'],4.);self.assertEqual(p.values['pitch'],1.)
        self.assertEqual(matcher.resolve(p.id).id,p.id)

    def test_price_does_not_transfer_between_models(self):
        from price_monitor.matching_ui import reference_price
        a,b=sensor('A'),sensor('B');matcher=Matcher([a,b])
        row={'our_article':'A','our_price':'123.45'}
        self.assertEqual(reference_price(row,evaluate(sensor(),a),matcher),'123.45')
        self.assertIsNone(reference_price(row,evaluate(sensor(),b),matcher))
        self.assertIsNone(reference_price({'our_price':'123.45'},evaluate(sensor(),a),matcher))

    def test_prices_do_not_enter_normalization(self):
        r={'model':'A','category':'Индуктивные','props':{'Размер резьбы корпуса':'M18x1'}}
        self.assertEqual(normalize_sensor({**r,'price':1}).values,normalize_sensor({**r,'price':100000}).values)

    def test_saved_choice_round_trip_and_excel_evidence(self):
        root=Path(__file__).resolve().parents[1]/'data';root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as tmp:
            store=Store(Path(tmp)/'prices.sqlite3')
            with store.connect() as c:c.execute('INSERT INTO rules VALUES(?,?,?,?,?,?,?,?)',(1,'1','megak','МЕГА-К','REF','https://mega-k.com/products/ref','','2026-09-29'))
            library=Library(store.catalog);library.select([1])
            library.save_comparisons([{'rule_id':1,'our_article':'ИВ05-NO-PNP(Л63)','our_price':'2500','our_currency':'RUB','note':'С оговоркой'}])
            _,rows=Library(store.catalog).products(selected=True)
            self.assertEqual(rows[0]['our_article'],'ИВ05-NO-PNP(Л63)');self.assertEqual(rows[0]['our_price'],'2500')
        match=evaluate(sensor(),sensor('Кандидат',sn=4.))
        wb=load_workbook(BytesIO(xlsx_bytes(matching_export(sensor(),match))))
        self.assertIn('Версия правил',[c.value for c in wb.active[1]])
        self.assertGreater(wb.active.max_row,15)

    def test_flowchart_is_valid_svg_and_has_explicit_outcomes(self):
        from price_monitor.algorithms import flowchart_svg, EDGES, NODES
        root=ET.fromstring(flowchart_svg());self.assertTrue(root.tag.endswith('svg'))
        ids={n[0] for n in NODES};self.assertTrue(all(a in ids and b in ids for a,b,_ in EDGES))
        self.assertTrue({'direct','close','unknown','reject'}<=ids)


if __name__=='__main__':unittest.main()
